from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Final

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FILL: Final = PatternFill("solid", fgColor="1F4E78")
SUBHEADER_FILL: Final = PatternFill("solid", fgColor="D9EAF7")
WHITE_FONT: Final = Font(color="FFFFFF", bold=True)

Row = dict[str, str]


def style_table_sheet(sheet: Worksheet, rows: list[Row], table_name: str) -> None:
    sheet.append(list(rows[0]))
    for row in rows:
        sheet.append(list(row.values()))
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    table = Table(displayName=table_name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    for index, column in enumerate(sheet.columns, start=1):
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)
        sheet.column_dimensions[get_column_letter(index)].width = width


def write_workbook(path: Path, tables: list[tuple[str, list[Row]]], readme: list[str]) -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "README"
    sheet.append([readme[0]])
    sheet.merge_cells("A1:B1")
    sheet["A1"].fill = HEADER_FILL
    sheet["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    for index in range(1, len(readme), 2):
        sheet.append([readme[index], readme[index + 1]])
        sheet.cell(sheet.max_row, 1).fill = SUBHEADER_FILL
        sheet.cell(sheet.max_row, 1).font = Font(bold=True)
    sheet.column_dimensions["A"].width = 23
    sheet.column_dimensions["B"].width = 88
    sheet.freeze_panes = "A2"
    for index, (name, rows) in enumerate(tables, start=1):
        assert isinstance(workbook.create_sheet(name), Worksheet)
        style_table_sheet(workbook.worksheets[-1], rows, f"CohortTable{index}")
    workbook.save(path)


def write_readme(path: Path, workbook_name: str, total: Row) -> None:
    text = f"""# SharePoint 업로드 안내

- 파일: `{workbook_name}`
- 권장 경로: `/03_Clinical_Study/연세세브란스/측정_Cohort/`
- 문서 버전: v1.0 (2026-07-14)
- 임상 레코드: 60명
- 독립 spectrum 피험자: {total["independent_spectrum_subjects"]}명
- 측정량: 1차 295 + 2차 재측정 295 = {total["total_spectrum_replicates"]} spectra

2차 데이터는 신규 피험자가 아니라 동일 검체 재측정이다. YNOR은 건강인만으로 구성된 군이 아니며 건강한 자, 췌장낭종, 기타 양성 질환, 만성 췌장염을 포함한다. `YNOR 21`은 2차 측정 당시 원라벨이며 1차 `YNOR 48`과 동일 검체다.

이 패키지는 직접식별자, 정확한 임상·측정 일자, 자유서술 과거력, 원본 파일 경로와 raw spectrum을 포함하지 않는다. 측정 시점은 월 단위로만 제공하며 원본 데이터는 통제된 WSL/임상 저장소에 유지한다.
"""
    _ = path.write_text(text, encoding="utf-8")


def write_archive(path: Path, package: Path, package_name: str, files: list[Path]) -> None:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for item in files:
            archive.write(item, Path(package_name) / item.relative_to(package))
