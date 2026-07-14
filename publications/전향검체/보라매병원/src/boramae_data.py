from __future__ import annotations

import csv
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Final

import numpy as np
import openpyxl
from scipy.signal import savgol_filter

REPO: Final = Path(__file__).resolve().parents[4]
sys.path.insert(0, str(REPO / "src"))

from sers.io import read_spectrum  # noqa: E402
from sers.preprocessing import calibrate_spectrum  # noqa: E402
from sers.signal import baseline_correction, snv  # noqa: E402

OUT: Final = REPO / "publications" / "전향검체" / "보라매병원"
FIG_DIR: Final = OUT / "figures"
TABLE_DIR: Final = OUT / "tables"
CLINICAL_XLSX: Final = REPO / "data" / "clinical_data" / "보라매 병원 임상정보.xlsx"
BORAMAE_ROOT: Final = REPO / "data" / "raw_data" / "20260709_BPRO,BNOR_1mW_0.05s_Ave100"
LEGACY_PRO_ROOT: Final = REPO / "data" / "raw_data" / "1. Prostate cancer (100개)"
CLEAN_MANIFEST: Final = REPO / "results" / "clean_cohort_20260605" / "clean_cohort_manifest.csv"
MODEL_GRID: Final = REPO / "artifacts" / "usersnet" / "v1.0.0" / "common_grid.npy"
LABELS: Final = ["Control", "Biopsy-negative", "Prostate cancer"]
SHORT_LABELS: Final = ["Control", "Biopsy-negative", "Prostate"]
COLORS: Final = {"Control": "#2C7FB8", "Biopsy-negative": "#7A5195", "Prostate cancer": "#D95F02"}
FILE_RE: Final = re.compile(r"^(BPRO|BNOR)\s+([0-9]+)_([0-9]+|ave)\.CSV$", re.IGNORECASE)


@dataclass(frozen=True, slots=True)
class ClinicalSample:
    label: str
    sample_no: str
    group: str
    grade_group: int | None
    excluded: bool


@dataclass(frozen=True, slots=True)
class SubjectSpectrum:
    sample: ClinicalSample
    mean_spectrum: np.ndarray
    replicate_spectra: np.ndarray


@dataclass(frozen=True, slots=True)
class CleanProSpectra:
    spectra: np.ndarray
    aligned_spectra: np.ndarray
    shifts: np.ndarray


def ensure_dirs() -> None:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    TABLE_DIR.mkdir(parents=True, exist_ok=True)


def row_is_excluded(sheet: openpyxl.worksheet.worksheet.Worksheet, row: int) -> bool:
    return any(
        sheet.cell(row, col).fill.fgColor.type == "rgb"
        and sheet.cell(row, col).fill.fgColor.rgb == "FFFFFF00"
        for col in range(1, sheet.max_column + 1)
    )


def parse_group(raw_group: str | None) -> str:
    mapped = {
        "control": "Control",
        "Elevated PSA, biopsy-negative (PSA↑/Bx−)": "Biopsy-negative",
        "prostate": "Prostate cancer",
        None: "Excluded",
    }.get(raw_group)
    if mapped is not None:
        return mapped
    msg = f"Unknown clinical group: {raw_group}"
    raise RuntimeError(msg)


def load_clinical_samples() -> list[ClinicalSample]:
    wb = openpyxl.load_workbook(CLINICAL_XLSX, data_only=True)
    sheet = wb.active
    headers = [sheet.cell(1, col).value for col in range(1, sheet.max_column + 1)]
    cols = {str(name): i + 1 for i, name in enumerate(headers) if name is not None}
    samples: list[ClinicalSample] = []
    for row in range(2, sheet.max_row + 1):
        raw_label = str(sheet.cell(row, cols["solum_label"]).value)
        grade_raw = sheet.cell(row, cols["Grade Group"]).value
        samples.append(
            ClinicalSample(
                label=raw_label,
                sample_no=raw_label.split()[-1],
                group=parse_group(sheet.cell(row, cols["group"]).value),
                grade_group=int(grade_raw) if isinstance(grade_raw, int) else None,
                excluded=row_is_excluded(sheet, row),
            )
        )
    return samples


def collect_boramae_files() -> dict[str, list[Path]]:
    files: dict[str, list[Path]] = {}
    for path in sorted(BORAMAE_ROOT.rglob("*.CSV")):
        match = FILE_RE.match(path.name)
        if match is None:
            continue
        group, sample_no, rep = match.groups()
        if rep.lower() != "ave":
            files.setdefault(f"{group.upper()} {sample_no}", []).append(path)
    return files


def raw_key(sample: ClinicalSample, files: dict[str, list[Path]]) -> str:
    for group in ("BPRO", "BNOR"):
        key = f"{group} {sample.sample_no}"
        if key in files:
            return key
    msg = f"No raw spectra found for {sample.label}"
    raise RuntimeError(msg)


def preprocess_arrays(
    x: np.ndarray, y: np.ndarray, grid: np.ndarray
) -> tuple[dict[str, tuple[np.ndarray, np.ndarray]], np.ndarray]:
    mask = (x >= float(grid.min())) & (x <= float(grid.max()))
    x_trim = x[mask]
    y_trim = y[mask]
    y_smooth = savgol_filter(y_trim, window_length=11, polyorder=3, mode="interp")
    y_base = baseline_correction(y_smooth, window=101)
    y_snv = snv(y_base)
    y_grid = np.interp(grid, x_trim, y_snv)
    return {
        "Raw": (x, y),
        "Trimmed": (x_trim, y_trim),
        "Smoothed": (x_trim, y_smooth),
        "Baseline corrected": (x_trim, y_base),
        "SNV + model grid": (grid, y_grid),
    }, y_grid


def preprocess_file(
    path: Path, grid: np.ndarray
) -> tuple[dict[str, tuple[np.ndarray, np.ndarray]], np.ndarray]:
    return preprocess_arrays(*read_spectrum(path), grid)


def preprocess_aligned_file(path: Path, grid: np.ndarray) -> tuple[np.ndarray, float]:
    x, y = read_spectrum(path)
    aligned_x, aligned_y, shift = calibrate_spectrum(x, y, window=20.0)
    return preprocess_arrays(aligned_x, aligned_y, grid)[1], shift


def build_boramae_subjects(
    samples: list[ClinicalSample], grid: np.ndarray
) -> list[SubjectSpectrum]:
    files = collect_boramae_files()
    subjects: list[SubjectSpectrum] = []
    for sample in samples:
        if sample.excluded or sample.group == "Excluded":
            continue
        spectra = [preprocess_file(path, grid)[1] for path in files[raw_key(sample, files)]]
        mat = np.vstack(spectra)
        subjects.append(
            SubjectSpectrum(sample=sample, mean_spectrum=mat.mean(axis=0), replicate_spectra=mat)
        )
    return subjects


def build_aligned_boramae_subjects(
    samples: list[ClinicalSample], grid: np.ndarray
) -> tuple[list[SubjectSpectrum], np.ndarray]:
    files = collect_boramae_files()
    subjects: list[SubjectSpectrum] = []
    shifts: list[np.ndarray] = []
    for sample in samples:
        if sample.excluded or sample.group == "Excluded":
            continue
        aligned = [preprocess_aligned_file(path, grid) for path in files[raw_key(sample, files)]]
        matrix = np.vstack([item[0] for item in aligned])
        shifts.append(np.asarray([item[1] for item in aligned], dtype=float))
        subjects.append(
            SubjectSpectrum(
                sample=sample, mean_spectrum=matrix.mean(axis=0), replicate_spectra=matrix
            )
        )
    return subjects, np.vstack(shifts)


def build_legacy_pro_subjects(grid: np.ndarray) -> np.ndarray:
    grouped: dict[str, list[Path]] = {}
    for path in sorted(LEGACY_PRO_ROOT.glob("PRO *.CSV")):
        match = re.match(r"^PRO\s+([0-9]+)_([0-9]+|ave)\.CSV$", path.name)
        if match is not None and match.group(2) != "ave":
            grouped.setdefault(match.group(1), []).append(path)
    rows = [
        np.vstack([preprocess_file(path, grid)[1] for path in paths]).mean(axis=0)
        for paths in grouped.values()
    ]
    return np.vstack(rows)


def clean_pro_ids() -> set[str]:
    with CLEAN_MANIFEST.open(encoding="utf-8-sig") as handle:
        rows = csv.DictReader(handle)
        return {str(int(float(row["sample_id"]))) for row in rows if row["source_group"] == "PRO"}


def build_clean_pro_alignment(grid: np.ndarray) -> CleanProSpectra:
    selected = clean_pro_ids()
    grouped: dict[str, list[Path]] = {}
    for path in sorted(LEGACY_PRO_ROOT.glob("PRO *.CSV")):
        match = re.match(r"^PRO\s+([0-9]+)_([0-9]+|ave)\.CSV$", path.name)
        if match is not None and match.group(1) in selected and match.group(2) != "ave":
            grouped.setdefault(match.group(1), []).append(path)
    ordered = sorted(grouped, key=int)
    spectra = [
        np.vstack([preprocess_file(path, grid)[1] for path in grouped[item]]).mean(axis=0)
        for item in ordered
    ]
    aligned_items = [
        [preprocess_aligned_file(path, grid) for path in grouped[item]] for item in ordered
    ]
    aligned = [np.vstack([result[0] for result in items]).mean(axis=0) for items in aligned_items]
    shifts = [result[1] for items in aligned_items for result in items]
    return CleanProSpectra(np.vstack(spectra), np.vstack(aligned), np.asarray(shifts, dtype=float))


def build_clean_pro_subjects(grid: np.ndarray) -> np.ndarray:
    return build_clean_pro_alignment(grid).spectra


def group_matrix(subjects: list[SubjectSpectrum], group: str) -> np.ndarray:
    return np.vstack([sub.mean_spectrum for sub in subjects if sub.sample.group == group])
