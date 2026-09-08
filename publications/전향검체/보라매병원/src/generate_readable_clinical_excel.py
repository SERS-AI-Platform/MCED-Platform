from __future__ import annotations

import csv
from pathlib import Path
from typing import Final

import openpyxl
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

REPO: Final = Path(__file__).resolve().parents[4]
OUT: Final = REPO / "publications" / "전향검체" / "보라매병원" / "internal"
SOURCE: Final = OUT / "clinical_classification_analysis_bpro.csv"
CLINICAL_XLSX: Final = REPO / "data" / "clinical_data" / "보라매 병원 임상정보.xlsx"
OUTPUT: Final = OUT / "clinical_classification_analysis_readable.xlsx"

FIELDS: Final = (
    ("BPRO 라벨", "case_id"),
    ("patient_code", "patient_code"),
    ("임상군", "publication_group"),
    ("모델 실제 라벨", "true_label"),
    ("PSA", "psa"),
    ("Gleason Score", "gleason_score"),
    ("Grade Group", "grade_group"),
    ("Grade band", "grade_band"),
    ("Three-group 예측", "three_group_pred_label"),
    ("Three-group 판정", "three_group_error_direction"),
    ("Three-group Cancer 확률", "three_group_prob_cancer"),
    ("Cancer vs Non-cancer 예측", "screening_pred_label"),
    ("Cancer vs Non-cancer 판정", "screening_error_direction"),
    ("Cancer 확률", "screening_prob_cancer"),
    ("Urine SG", "ua_sg"),
    ("Urine pH", "ua_ph"),
    ("Urine protein", "ua_protein"),
    ("Urine blood", "ua_blood"),
    ("Urine leukocyte", "ua_leukocyte"),
    ("Creatinine", "creatinine"),
    ("BUN", "bun"),
    ("Glucose", "glucose"),
    ("HbA1c", "hba1c"),
)

HEADER_FILL: Final = PatternFill("solid", fgColor="1F4E78")
SUBHEADER_FILL: Final = PatternFill("solid", fgColor="D9EAF7")
ERROR_FILL: Final = PatternFill("solid", fgColor="FCE4D6")
HEADER_FONT: Final = Font(color="FFFFFF", bold=True)


def read_rows() -> list[dict[str, str]]:
    with SOURCE.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def read_patient_codes() -> dict[str, str]:
    sheet = openpyxl.load_workbook(CLINICAL_XLSX, data_only=True, read_only=True).active
    rows = list(sheet.values)
    return {
        str(row[1]): str(row[3]) for row in rows[1:] if row[1] is not None and row[3] is not None
    }


def normalize_group(group: str) -> str:
    return "Cancer" if group == "Prostate cancer" else group


def value(row: dict[str, str], field: str) -> str | float | int:
    if field == "patient_code":
        return ""
    raw = row[field]
    if raw == "":
        return ""
    if field in {"psa", "ua_sg", "ua_ph", "creatinine", "bun", "glucose", "hba1c"}:
        return float(raw)
    if field in {"grade_group", "patient_code"}:
        return int(float(raw))
    if field in {"three_group_prob_cancer", "screening_prob_cancer"}:
        return round(float(raw), 4)
    return raw


def add_analysis_sheet(
    workbook: Workbook, rows: list[dict[str, str]], patient_codes: dict[str, str]
) -> None:
    sheet = workbook.active
    sheet.title = "분류결과_전체"
    headers = [label for label, _ in FIELDS]
    sheet.append(headers)
    for header_cell in sheet[1]:
        header_cell.fill = HEADER_FILL
        header_cell.font = HEADER_FONT
        header_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        values = [value(row, field) for _, field in FIELDS]
        values[1] = patient_codes[row["case_id"]]
        values[2] = normalize_group(str(values[2]))
        values[9] = "Correct" if row["three_group_is_correct"] == "True" else "Incorrect"
        values[12] = "Correct" if row["screening_is_correct"] == "True" else "Incorrect"
        sheet.append(values)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 36
    widths = [
        14,
        14,
        18,
        16,
        10,
        17,
        12,
        12,
        20,
        28,
        18,
        18,
        24,
        18,
        12,
        10,
        16,
        14,
        16,
        14,
        10,
        12,
        12,
    ]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    table = Table(displayName="ClinicalClassification", ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
    )
    sheet.add_table(table)
    result_columns = {"J", "M"}
    for column in result_columns:
        sheet.conditional_formatting.add(
            f"{column}2:{column}{sheet.max_row}",
            FormulaRule(formula=[f'NOT(OR({column}2="correct",{column}2=""))'], fill=ERROR_FILL),
        )


def add_error_sheet(
    workbook: Workbook, rows: list[dict[str, str]], patient_codes: dict[str, str]
) -> None:
    sheet = workbook.create_sheet("오분류만")
    fields = (
        "BPRO 라벨",
        "patient_code",
        "임상군",
        "Grade Group",
        "Gleason Score",
        "Three-group 예측",
        "Three-group 오분류 방향",
        "Cancer vs Non-cancer 예측",
        "Cancer vs Non-cancer 오분류 방향",
    )
    keys = (
        "case_id",
        "publication_group",
        "grade_group",
        "gleason_score",
        "three_group_pred_label",
        "three_group_error_direction",
        "screening_pred_label",
        "screening_error_direction",
    )
    sheet.append(fields)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        if row["three_group_is_correct"] == "True" and row["screening_is_correct"] == "True":
            continue
        values = [value(row, key) for key in keys]
        values.insert(1, patient_codes[row["case_id"]])
        values[2] = normalize_group(str(values[2]))
        sheet.append(values)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate((14, 14, 18, 12, 17, 20, 30, 24, 30), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.fill = ERROR_FILL
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def add_legend_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("항목설명")
    rows = [
        ("Three-group", "Control / Biopsy-negative / Cancer 분류"),
        ("Cancer vs Non-cancer", "Cancer와 Non-cancer를 구분하는 이진 분류"),
        ("correct", "실제 라벨과 예측 라벨이 일치"),
        ("cancer_to_biopsy_negative", "실제 Cancer를 Biopsy-negative로 예측"),
        ("cancer_to_control", "실제 Cancer를 Control로 예측"),
        ("cancer_to_non_cancer", "실제 Cancer를 Non-cancer로 예측"),
        ("control_to_biopsy_negative", "실제 Control을 Biopsy-negative로 예측"),
        ("biopsy_negative_to_control", "실제 Biopsy-negative를 Control로 예측"),
        ("non_cancer_to_cancer", "실제 Non-cancer를 Cancer로 예측"),
        ("주의", "BPRO 라벨과 patient_code는 내부 매칭용이며 외부 공개용이 아님"),
    ]
    sheet.append(("항목", "설명"))
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 80
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def main() -> None:
    rows = read_rows()
    patient_codes = read_patient_codes()
    workbook = Workbook()
    add_analysis_sheet(workbook, rows, patient_codes)
    add_error_sheet(workbook, rows, patient_codes)
    add_legend_sheet(workbook)
    workbook.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
