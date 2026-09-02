from __future__ import annotations

import re
from collections.abc import Mapping
from contextlib import closing
from dataclasses import dataclass
from pathlib import Path
from typing import Final

import numpy as np
import openpyxl
from boramae_data import preprocess_file

from powder_comparison.cohort import PairedSubject, SubjectFiles

LABEL_PATTERN: Final = re.compile(r"(?P<sample_id>\d+)$")
CLINICAL_GROUPS: Final = {
    "control": "Control",
    "Elevated PSA, biopsy-negative (PSA↑/Bx−)": "Biopsy-negative",
    "prostate": "Cancer",
}


class ClinicalDataError(RuntimeError):
    pass


@dataclass(frozen=True, slots=True)
class PreprocessingStageMatrices:
    raw: np.ndarray
    smoothed: np.ndarray
    baseline_corrected: np.ndarray
    snv: np.ndarray


@dataclass(frozen=True, slots=True)
class ModalitySpectra:
    spectra: np.ndarray
    raw_mean_intensity: np.ndarray
    raw_max_intensity: np.ndarray
    replicate_correlation: np.ndarray
    stages: PreprocessingStageMatrices
    replicate_stages: PreprocessingStageMatrices | None = None


@dataclass(frozen=True, slots=True)
class PairedSpectra:
    grid: np.ndarray
    sample_ids: np.ndarray
    clinical_groups: np.ndarray
    prefixes: np.ndarray
    powder_order: np.ndarray
    legacy: ModalitySpectra
    powder: ModalitySpectra


@dataclass(frozen=True, slots=True)
class ClinicalCovariates:
    patient_ids: np.ndarray
    psa: np.ndarray
    psa_bands: np.ndarray


@dataclass(frozen=True, slots=True)
class _ProcessedSubject:
    spectrum: np.ndarray
    raw_mean: float
    raw_maximum: float
    replicate_correlation: float
    stages: tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]
    replicate_stages: tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]


def _clinical_group(raw_group: str | None) -> str:
    if raw_group is None:
        return "Excluded"
    try:
        return CLINICAL_GROUPS[raw_group]
    except KeyError as error:
        raise ClinicalDataError(f"Unknown clinical group: {raw_group}") from error


def load_clinical_labels(path: Path) -> Mapping[int, str]:
    with closing(openpyxl.load_workbook(path, read_only=True, data_only=True)) as workbook:
        sheet = workbook.worksheets[0]
        headers = {
            str(sheet.cell(1, column).value): column for column in range(1, sheet.max_column + 1)
        }
        labels: dict[int, str] = {}
        for row in range(2, sheet.max_row + 1):
            raw_label = str(sheet.cell(row, headers["solum_label"]).value)
            match = LABEL_PATTERN.search(raw_label)
            if match is None:
                raise ClinicalDataError(f"Missing sample number in clinical label: {raw_label}")
            raw_group_value = sheet.cell(row, headers["group"]).value
            raw_group = None if raw_group_value is None else str(raw_group_value)
            group = _clinical_group(raw_group)
            if group != "Excluded":
                labels[int(match["sample_id"])] = group
    return labels


def _psa_band(value: float) -> str:
    if not np.isfinite(value):
        return "Missing"
    if value < 4.0:
        return "<4"
    if value < 10.0:
        return "4-<10"
    return ">=10"


def load_clinical_covariates(path: Path, sample_ids: np.ndarray) -> ClinicalCovariates:
    with closing(openpyxl.load_workbook(path, read_only=True, data_only=True)) as workbook:
        sheet = workbook.worksheets[0]
        headers = {
            str(sheet.cell(1, column).value): column for column in range(1, sheet.max_column + 1)
        }
        patient_id_by_sample: dict[int, str] = {}
        psa_by_sample: dict[int, float] = {}
        for row in range(2, sheet.max_row + 1):
            raw_label = str(sheet.cell(row, headers["solum_label"]).value)
            match = LABEL_PATTERN.search(raw_label)
            if match is None:
                raise ClinicalDataError(f"Missing sample number in clinical label: {raw_label}")
            sample_id = int(match["sample_id"])
            raw_patient_id = sheet.cell(row, headers["patient_code"]).value
            if raw_patient_id is None:
                raise ClinicalDataError(f"Missing patient code for clinical label: {raw_label}")
            patient_id_by_sample[sample_id] = str(int(float(str(raw_patient_id))))
            raw_psa = sheet.cell(row, headers["psa"]).value
            psa_by_sample[sample_id] = np.nan if raw_psa is None else float(str(raw_psa))
    patient_ids = np.array([patient_id_by_sample[int(sample_id)] for sample_id in sample_ids])
    psa = np.array([psa_by_sample[int(sample_id)] for sample_id in sample_ids], dtype=float)
    return ClinicalCovariates(
        patient_ids=patient_ids,
        psa=psa,
        psa_bands=np.array([_psa_band(value) for value in psa]),
    )


def _replicate_correlation(matrix: np.ndarray) -> float:
    correlations = np.corrcoef(matrix)
    upper = correlations[np.triu_indices(len(matrix), k=1)]
    return float(np.mean(upper))


def _process_subject(subject: SubjectFiles, grid: np.ndarray) -> _ProcessedSubject:
    processed: list[np.ndarray] = []
    raw_means: list[float] = []
    raw_maxima: list[float] = []
    stage_rows: list[list[np.ndarray]] = [[], [], [], []]
    for path in subject.replicates:
        stages, processed_spectrum = preprocess_file(path, grid)
        raw_x, raw_intensity = stages["Raw"]
        processed.append(processed_spectrum)
        raw_means.append(float(np.mean(raw_intensity)))
        raw_maxima.append(float(np.max(raw_intensity)))
        for rows, name in zip(
            stage_rows,
            ("Raw", "Smoothed", "Baseline corrected", "SNV + model grid"),
            strict=True,
        ):
            stage_x, stage_y = stages[name]
            rows.append(np.interp(grid, stage_x if name != "Raw" else raw_x, stage_y))
    matrix = np.vstack(processed)
    return _ProcessedSubject(
        spectrum=matrix.mean(axis=0),
        raw_mean=float(np.mean(raw_means)),
        raw_maximum=float(np.mean(raw_maxima)),
        replicate_correlation=_replicate_correlation(matrix),
        stages=(
            np.vstack(stage_rows[0]).mean(axis=0),
            np.vstack(stage_rows[1]).mean(axis=0),
            np.vstack(stage_rows[2]).mean(axis=0),
            np.vstack(stage_rows[3]).mean(axis=0),
        ),
        replicate_stages=(
            np.vstack(stage_rows[0]),
            np.vstack(stage_rows[1]),
            np.vstack(stage_rows[2]),
            np.vstack(stage_rows[3]),
        ),
    )


def _modality_spectra(
    pairs: tuple[PairedSubject, ...], grid: np.ndarray, *, powder: bool
) -> ModalitySpectra:
    rows = [_process_subject(pair.powder if powder else pair.legacy, grid) for pair in pairs]
    return ModalitySpectra(
        spectra=np.vstack([row.spectrum for row in rows]),
        raw_mean_intensity=np.array([row.raw_mean for row in rows], dtype=float),
        raw_max_intensity=np.array([row.raw_maximum for row in rows], dtype=float),
        replicate_correlation=np.array([row.replicate_correlation for row in rows], dtype=float),
        stages=PreprocessingStageMatrices(
            raw=np.vstack([row.stages[0] for row in rows]),
            smoothed=np.vstack([row.stages[1] for row in rows]),
            baseline_corrected=np.vstack([row.stages[2] for row in rows]),
            snv=np.vstack([row.stages[3] for row in rows]),
        ),
        replicate_stages=PreprocessingStageMatrices(
            raw=np.stack([row.replicate_stages[0] for row in rows]),
            smoothed=np.stack([row.replicate_stages[1] for row in rows]),
            baseline_corrected=np.stack([row.replicate_stages[2] for row in rows]),
            snv=np.stack([row.replicate_stages[3] for row in rows]),
        ),
    )


def build_paired_spectra(pairs: tuple[PairedSubject, ...], grid: np.ndarray) -> PairedSpectra:
    return PairedSpectra(
        grid=grid.copy(),
        sample_ids=np.array([pair.sample_id for pair in pairs], dtype=int),
        clinical_groups=np.array([pair.clinical_group for pair in pairs]),
        prefixes=np.array([pair.powder.prefix for pair in pairs]),
        powder_order=np.array([pair.powder.order_index for pair in pairs], dtype=int),
        legacy=_modality_spectra(pairs, grid, powder=False),
        powder=_modality_spectra(pairs, grid, powder=True),
    )
