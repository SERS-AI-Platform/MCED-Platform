from __future__ import annotations

import csv
import re
from contextlib import closing
from dataclasses import dataclass
from pathlib import Path
from typing import Final

import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from matplotlib.figure import Figure
from matplotlib.font_manager import FontProperties
from scipy.signal import savgol_filter

from sers.io import read_spectrum
from sers.signal import baseline_correction, snv

PRO_PATTERN: Final = re.compile(r"^PRO\s+(?P<sample_id>\d+)_(?P<replicate>\d+|ave)\.CSV$")
BORAMAE_PATTERN: Final = re.compile(
    r"^(?:BPRO|BNOR)\s+(?P<sample_id>\d+)_(?P<replicate>\d+|ave)\.CSV$"
)


class CohortDataError(RuntimeError):
    pass


@dataclass(frozen=True, slots=True)
class CohortSeries:
    label: str
    raw_spectra: np.ndarray
    processed_spectra: np.ndarray
    color: str
    line_style: str


def _read_spectrum_pair(path: Path, grid: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    x, y = read_spectrum(path)
    mask = (x >= float(grid.min())) & (x <= float(grid.max()))
    x_trim = x[mask]
    y_trim = y[mask]
    smoothed = np.asarray(
        savgol_filter(y_trim, window_length=11, polyorder=3, mode="interp"),
        dtype=float,
    )
    corrected = baseline_correction(smoothed, window=101)
    return np.interp(grid, x_trim, y_trim), np.interp(grid, x_trim, snv(corrected))


def _subject_means(
    grouped: dict[int, list[Path]], grid: np.ndarray
) -> tuple[np.ndarray, np.ndarray]:
    raw_means: list[np.ndarray] = []
    processed_means: list[np.ndarray] = []
    for sample_id in sorted(grouped):
        pairs = [_read_spectrum_pair(path, grid) for path in grouped[sample_id]]
        raw_rows, processed_rows = zip(*pairs, strict=True)
        raw_means.append(np.vstack(raw_rows).mean(axis=0))
        processed_means.append(np.vstack(processed_rows).mean(axis=0))
    return np.vstack(raw_means), np.vstack(processed_means)


def _clean_pro_spectra(repo: Path, grid: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    manifest = repo / "results" / "clean_cohort_20260605" / "clean_cohort_manifest.csv"
    with manifest.open(encoding="utf-8-sig") as handle:
        selected = {
            int(float(row["sample_id"]))
            for row in csv.DictReader(handle)
            if row["source_group"] == "PRO"
        }
    root = repo / "data" / "raw_data" / "1. Prostate cancer (100개)"
    grouped: dict[int, list[Path]] = {}
    for path in sorted(root.glob("PRO *.CSV")):
        match = PRO_PATTERN.fullmatch(path.name)
        if match is None or match["replicate"].lower() == "ave":
            continue
        sample_id = int(match["sample_id"])
        if sample_id in selected:
            grouped.setdefault(sample_id, []).append(path)
    return _subject_means(grouped, grid)


def _boramae_cancer_ids(repo: Path) -> set[int]:
    path = repo / "data" / "clinical_data" / "보라매 병원 임상정보.xlsx"
    with closing(openpyxl.load_workbook(path, data_only=True)) as workbook:
        sheet = workbook.active
        if sheet is None:
            raise CohortDataError(f"Clinical workbook has no active worksheet: {path}")
        headers = {
            str(sheet.cell(1, column).value): column
            for column in range(1, sheet.max_column + 1)
        }
        selected: set[int] = set()
        for row in range(2, sheet.max_row + 1):
            is_excluded = any(
                sheet.cell(row, column).fill.fgColor.type == "rgb"
                and sheet.cell(row, column).fill.fgColor.rgb == "FFFFFF00"
                for column in range(1, sheet.max_column + 1)
            )
            if sheet.cell(row, headers["group"]).value == "prostate" and not is_excluded:
                label = str(sheet.cell(row, headers["solum_label"]).value)
                selected.add(int(label.split()[-1]))
    return selected


def _boramae_cancer_spectra(
    repo: Path, grid: np.ndarray
) -> tuple[np.ndarray, np.ndarray]:
    selected = _boramae_cancer_ids(repo)
    root = repo / "data" / "raw_data" / "20260709_BPRO,BNOR_1mW_0.05s_Ave100"
    grouped: dict[int, list[Path]] = {}
    for path in sorted(root.rglob("*.CSV")):
        match = BORAMAE_PATTERN.fullmatch(path.name)
        if match is None or match["replicate"].lower() == "ave":
            continue
        sample_id = int(match["sample_id"])
        if sample_id in selected:
            grouped.setdefault(sample_id, []).append(path)
    return _subject_means(grouped, grid)


def _mapping_spectra(repo: Path, grid: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    root = repo / "data" / "mapping"
    with closing(
        openpyxl.load_workbook(root / "clinical_df.xlsx", read_only=True, data_only=True)
    ) as workbook:
        sheet = workbook.active
        if sheet is None:
            raise CohortDataError("Mapping clinical workbook has no active worksheet")
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        columns = {str(value): index for index, value in enumerate(header) if value is not None}
        selected = {
            str(row[columns["solum_label"]]).replace("_", " ")
            for row in rows
            if row[columns["cohort_group"]] == "prostate"
        }
    folders = {
        folder.name: folder
        for date_folder in root.glob("2026*_mapping")
        for folder in date_folder.iterdir()
        if folder.is_dir()
    }
    grouped = {
        ordinal: [
            path
            for path in sorted(folders[label].glob(f"{label}_*.CSV"))
            if (match := BORAMAE_PATTERN.fullmatch(path.name)) is not None
            and match["replicate"].lower() != "ave"
        ]
        for ordinal, label in enumerate(sorted(selected), start=1)
    }
    return _subject_means(grouped, grid)


def load_three_cohort_series(
    repo: Path, grid: np.ndarray
) -> tuple[CohortSeries, CohortSeries, CohortSeries]:
    clean_raw, clean_processed = _clean_pro_spectra(repo, grid)
    liquid_raw, liquid_processed = _boramae_cancer_spectra(repo, grid)
    powder_raw, powder_processed = _mapping_spectra(repo, grid)
    cohorts = (
        CohortSeries("후향 전립선암 데이터 (n=91)", clean_raw, clean_processed, "#4D4D4D", "-"),
        CohortSeries(
            "보라매 검체 (액상 환원제, n=41)",
            liquid_raw,
            liquid_processed,
            "#D95F02",
            "--",
        ),
        CohortSeries(
            "보라매 검체 (분말 환원제, n=43)",
            powder_raw,
            powder_processed,
            "#2C7FB8",
            "-.",
        ),
    )
    counts = tuple(len(cohort.processed_spectra) for cohort in cohorts)
    if counts != (91, 41, 43):
        raise CohortDataError(f"Unexpected cohort counts: {counts}; expected (91, 41, 43)")
    if any(
        series.shape[1] != len(grid)
        for cohort in cohorts
        for series in (cohort.raw_spectra, cohort.processed_spectra)
    ):
        raise CohortDataError("Cohort spectra do not share the canonical grid")
    return cohorts


def create_three_cohort_figure(
    grid: np.ndarray, cohorts: tuple[CohortSeries, ...]
) -> Figure:
    figure, axes = plt.subplots(
        2, 1, figsize=(12, 8.4), sharex=True, layout="constrained"
    )
    for axis, arrays in (
        (axes[0], tuple(cohort.raw_spectra for cohort in cohorts)),
        (axes[1], tuple(cohort.processed_spectra for cohort in cohorts)),
    ):
        for cohort, spectra in zip(cohorts, arrays, strict=True):
            mean = spectra.mean(axis=0)
            sem = spectra.std(axis=0, ddof=1) / np.sqrt(len(spectra))
            axis.plot(
                grid,
                mean,
                label=cohort.label,
                color=cohort.color,
                linestyle=cohort.line_style,
                linewidth=2.0,
            )
            axis.fill_between(
                grid,
                mean - 1.96 * sem,
                mean + 1.96 * sem,
                color=cohort.color,
                alpha=0.12,
                linewidth=0,
            )
        axis.grid(axis="y", color="#D9D9D9", linewidth=0.8, alpha=0.7)
        axis.spines[["top", "right"]].set_visible(False)
    figure.suptitle(
        "Prostate cancer spectra before and after preprocessing",
        x=0.065,
        ha="left",
        fontsize=15,
    )
    axes[0].set_title(
        "A  Trimmed raw spectra: subject means with 95% CI; common-grid interpolation only",
        loc="left",
        color="#555555",
        fontsize=9,
        pad=12,
    )
    axes[1].set_title(
        f"B  Processed spectra: subject means with 95% CI on one {len(grid)}-point grid",
        loc="left",
        color="#555555",
        fontsize=9,
        pad=12,
    )
    axes[0].set_ylabel("Raw intensity (a.u.)")
    axes[1].set_ylabel("SNV intensity")
    axes[1].set_xlabel("Raman shift (cm$^{-1}$)")
    axes[0].legend(
        frameon=False,
        loc="upper right",
        prop=FontProperties(family="Noto Sans CJK KR"),
    )
    return figure


def write_three_cohort_figure(repo: Path) -> Path:
    grid = np.load(repo / "artifacts" / "usersnet" / "v1.0.0" / "common_grid.npy")
    cohorts = load_three_cohort_series(repo, grid)
    figure = create_three_cohort_figure(grid, cohorts)
    output = (
        repo
        / "publications"
        / "전향검체"
        / "보라매병원"
        / "figures"
        / "fig02_boramae_vs_legacy_pro_processed"
    )
    figure.savefig(output.with_suffix(".png"), dpi=300, bbox_inches="tight")
    figure.savefig(output.with_suffix(".pdf"), bbox_inches="tight")
    plt.close(figure)
    return output.with_suffix(".png")
