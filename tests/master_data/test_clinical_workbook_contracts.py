from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from sers.master_data.clinical_ingestion import (
    ClinicalIngestionRequest,
    ingest_clinical_source,
)
from sers.master_data.clinical_inventory import inventory_clinical_sources
from sers.master_data.clinical_parser import (
    MissingClinicalSheetError,
    load_clinical_rows,
)
from tests.master_data.clinical_test_support import database


def test_source_specific_workbook_shapes_preserve_all_event_rows(
    tmp_path: Path,
) -> None:
    # Given: synthetic workbooks matching the reviewed CRC, BLA, LUN2, and LUN3 shapes.
    crc_path = tmp_path / "9. 대장암/SMCXD06_대장암.xlsx"
    crc_path.parent.mkdir(parents=True)
    crc = Workbook()
    early = crc.active
    early.title = "대장암1~2기270명"
    early.append(("NO", "제공자:제공자bCODE", "stage"))
    early.append((1, "CRC-SYNTH-1", "early"))
    late = crc.create_sheet("대장암3~4기30명")
    late.append(("NO", "제공자:제공자bCODE", "stage"))
    late.append((1, "CRC-SYNTH-2", "late"))
    crc.save(crc_path)

    bladder_path = tmp_path / "11. 방광암/SMCXD06_방광암.xlsm"
    bladder_path.parent.mkdir(parents=True)
    bladder = Workbook()
    bladder_sheet = bladder.active
    bladder_sheet.title = "분양명단"
    bladder_sheet.append(("title",))
    bladder_sheet.append(("NO", "pathology"))
    bladder_sheet.append(("BLA-SYNTH-1", "synthetic"))
    bladder.save(bladder_path)

    lung2_path = tmp_path / "4. 폐암/SMCXD06_폐암 2.xlsx"
    lung2_path.parent.mkdir(parents=True)
    lung2 = Workbook()
    lung2_sheet = lung2.active
    lung2_sheet.title = "Sheet1"
    lung2_sheet.append(("no", "stage"))
    lung2_sheet.append(("LUN-SYNTH-2", "synthetic"))
    lung2.save(lung2_path)

    lung3_path = tmp_path / "4. 폐암/SMCXD06_폐암 3.xlsx"
    lung3 = Workbook()
    resource = lung3.active
    resource.title = "자원리스트"
    resource.append((None, "제공자:제공자bCODE", "collection_date"))
    resource.append(("LUN-SYNTH-3", "BCODE-SYNTH-3", "2026-01-01"))
    for sheet_name, field_name in (
        ("인구학적정보 및 암 관련 정보", "age"),
        ("병리검사", "Histologic"),
        ("진단검사", "WBC"),
    ):
        sheet = lung3.create_sheet(sheet_name)
        sheet.append(("제공자:제공자bCODE", field_name))
        sheet.append(("BCODE-SYNTH-3", "synthetic"))
    lung3.save(lung3_path)

    # When: the built-in registry contracts parse each physical workbook.
    inventory = inventory_clinical_sources(tmp_path)
    sources = {source.path.name: source for source in inventory.sources}
    rows = {
        name: load_clinical_rows(sources[name])
        for name in (
            "SMCXD06_대장암.xlsx",
            "SMCXD06_방광암.xlsm",
            "SMCXD06_폐암 2.xlsx",
            "SMCXD06_폐암 3.xlsx",
        )
    }

    # Then: all configured sheets become distinct provenance-bearing events.
    assert tuple(row.locator for row in rows["SMCXD06_대장암.xlsx"]) == (
        "대장암1~2기270명!row:2",
        "대장암3~4기30명!row:2",
    )
    assert tuple(row.locator for row in rows["SMCXD06_방광암.xlsm"]) == (
        "분양명단!row:3",
    )
    assert rows["SMCXD06_폐암 2.xlsx"][0].patient_id == "LUN-SYNTH-2"
    lung3_rows = rows["SMCXD06_폐암 3.xlsx"]
    assert tuple(row.event_type for row in lung3_rows) == (
        "collection",
        "other",
        "diagnosis",
        "lab",
    )
    assert {row.patient_id for row in lung3_rows} == {"LUN-SYNTH-3"}
    assert len({row.locator for row in lung3_rows}) == 4


def test_partial_multisheet_workbook_fails_atomically(tmp_path: Path) -> None:
    # Given: a LUN3 workbook missing the configured diagnostic-lab sheet.
    source_path = tmp_path / "4. 폐암/SMCXD06_폐암 3.xlsx"
    source_path.parent.mkdir(parents=True)
    workbook = Workbook()
    resource = workbook.active
    resource.title = "자원리스트"
    resource.append((None, "제공자:제공자bCODE"))
    resource.append(("LUN-SYNTH-4", "BCODE-SYNTH-4"))
    for sheet_name in ("인구학적정보 및 암 관련 정보", "병리검사"):
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(("제공자:제공자bCODE", "synthetic_field"))
        sheet.append(("BCODE-SYNTH-4", "synthetic"))
    workbook.save(source_path)
    contract = inventory_clinical_sources(tmp_path).sources[0]

    # When: the incomplete physical workbook crosses the ingestion boundary.
    with database(tmp_path / "clinical.db") as connection:
        with pytest.raises(MissingClinicalSheetError):
            ingest_clinical_source(
                connection,
                ClinicalIngestionRequest(contract, tmp_path / "raw"),
            )

        # Then: no partial sheet, source asset, or raw blob is published.
        assert connection.execute("SELECT COUNT(*) FROM source_assets").fetchone()[0] == 0
        assert connection.execute("SELECT COUNT(*) FROM clinical_events").fetchone()[0] == 0
        assert not tuple((tmp_path / "raw").rglob("*"))
