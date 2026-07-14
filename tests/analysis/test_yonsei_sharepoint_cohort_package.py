from __future__ import annotations

import csv
import re
import sys
import zipfile
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[2] / "publications" / "전향검체" / "연세세브란스 병원"
PACKAGE = ROOT / "sharepoint_upload" / "2026_SERS-AI_Yonsei_Prospective_Cohort_v1.0"
WORKBOOK = PACKAGE / "2026_SERS-AI_Yonsei_Prospective_Cohort_v1.0.xlsx"
ARCHIVE = PACKAGE.parent / "2026_SERS-AI_Yonsei_Prospective_Cohort_v1.0.zip"
SRC = ROOT / "src"
sys.path.insert(0, str(SRC))
PACKAGE_AVAILABLE = pytest.mark.skipif(
    not WORKBOOK.exists() or not ARCHIVE.exists(),
    reason="SharePoint acceptance package is stored outside Git",
)

import generate_sharepoint_cohort_package as generator  # noqa: E402
from yonsei_sharepoint.outputs import write_workbook  # noqa: E402

GENERATOR_INPUTS_AVAILABLE = pytest.mark.skipif(
    not any(generator.REACQUIRED_ROOT.rglob("*.CSV")),
    reason="SharePoint regeneration requires external reacquired spectra",
)


def read_csv(name: str) -> list[dict[str, str]]:
    with (PACKAGE / "csv" / name).open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def test_workbook_writer_serializes_core_properties(tmp_path: Path) -> None:
    path = tmp_path / "cohort.xlsx"

    write_workbook(path, [("Summary", [{"group": "YNOR", "count": "29"}])], ["Title"])

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    assert workbook.sheetnames == ["README", "Summary"]


def configure_outputs(monkeypatch: pytest.MonkeyPatch, root: Path) -> None:
    package = root / "package"
    monkeypatch.setattr(generator, "PACKAGE", package)
    monkeypatch.setattr(generator, "CSV_DIR", package / "csv")
    monkeypatch.setattr(generator, "WORKBOOK", package / "cohort.xlsx")
    monkeypatch.setattr(generator, "ARCHIVE", root / "cohort.zip")


def test_generator_rejects_missing_reacquired_measurements(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    empty = tmp_path / "empty"
    empty.mkdir()
    configure_outputs(monkeypatch, tmp_path)
    monkeypatch.setattr(generator, "REACQUIRED_ROOT", empty)

    with pytest.raises(RuntimeError, match="reacquired measurement inventory"):
        generator.main()


@GENERATOR_INPUTS_AVAILABLE
def test_generator_removes_stale_package_files(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    configure_outputs(monkeypatch, tmp_path)
    generator.PACKAGE.mkdir(parents=True)
    stale = generator.PACKAGE / "stale-private-export.txt"
    _ = stale.write_text("private", encoding="utf-8")

    generator.main()

    assert not stale.exists()
    with zipfile.ZipFile(generator.ARCHIVE) as archive:
        assert not any("stale-private-export.txt" in name for name in archive.namelist())


@PACKAGE_AVAILABLE
def test_workbook_contains_review_ready_sheets() -> None:
    workbook = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)

    assert workbook.sheetnames == [
        "README",
        "Cohort_Summary",
        "Measurement_Schedule",
        "Subject_Manifest",
        "Acquisition_Inventory",
        "YNOR_ID_Mapping",
        "Clinical_Completeness",
    ]


@PACKAGE_AVAILABLE
def test_package_counts_independent_subjects_and_acquisitions() -> None:
    summary = {row["source_group"]: row for row in read_csv("01_cohort_summary.csv")}
    acquisitions = read_csv("04_acquisition_inventory.csv")

    assert summary["TOTAL"]["clinical_records"] == "60"
    assert summary["TOTAL"]["independent_spectrum_subjects"] == "59"
    assert summary["TOTAL"]["total_spectrum_replicates"] == "590"
    assert len(acquisitions) == 118
    assert sum(row["acquisition_round"] == "1차" for row in acquisitions) == 59
    assert sum(row["acquisition_round"] == "2차 재측정" for row in acquisitions) == 59
    assert (
        sum(int(row["n_replicates"]) for row in acquisitions if row["acquisition_round"] == "1차")
        == 295
    )
    assert (
        sum(
            int(row["n_replicates"])
            for row in acquisitions
            if row["acquisition_round"] == "2차 재측정"
        )
        == 295
    )
    assert all(
        not row["measurement_month"] for row in acquisitions if row["acquisition_round"] == "1차"
    )
    assert all(
        row["measurement_month"] for row in acquisitions if row["acquisition_round"] == "2차 재측정"
    )


@PACKAGE_AVAILABLE
def test_subject_manifest_is_minimized_and_pseudonymized() -> None:
    rows = read_csv("03_subject_manifest.csv")
    forbidden = {
        "patient_id",
        "patient_code",
        "source_file",
        "past_history",
        "diagnosis_date",
        "sample_date",
        "surgery_date",
        "chemo_date",
        "age",
        "bmi",
    }

    assert len(rows) == 60
    assert not forbidden.intersection(rows[0])
    assert sum(row["spectrum_status"] == "임상정보만 존재" for row in rows) == 1
    assert all(not value.startswith("/home/") for row in rows for value in row.values())


@PACKAGE_AVAILABLE
def test_sharepoint_tables_exclude_exact_dates() -> None:
    exact_date = re.compile(r"\b\d{4}-\d{2}-\d{2}\b")
    rows = [row for path in sorted((PACKAGE / "csv").glob("*.csv")) for row in read_csv(path.name)]

    assert not any(exact_date.search(value) for row in rows for value in row.values())
    assert not any("measurement_date" in row for row in rows)


@PACKAGE_AVAILABLE
def test_csv_files_have_excel_compatible_bom() -> None:
    csv_files = sorted((PACKAGE / "csv").glob("*.csv"))

    assert len(csv_files) == 6
    assert all(path.read_bytes().startswith(b"\xef\xbb\xbf") for path in csv_files)


@PACKAGE_AVAILABLE
def test_ynor_repeat_label_maps_to_primary_subject() -> None:
    rows = read_csv("05_ynor_id_mapping.csv")
    mapped = next(row for row in rows if row["original_reacquired_label"] == "YNOR 21")

    assert mapped["subject_case_id"] == "YON-C-021"
    assert mapped["primary_label"] == "YNOR 48"
    assert mapped["current_reacquired_label"] == "YNOR 48"


@PACKAGE_AVAILABLE
def test_archive_contains_only_sharepoint_deliverables() -> None:
    with zipfile.ZipFile(ARCHIVE) as archive:
        names = set(archive.namelist())

    prefix = f"{PACKAGE.name}/"
    expected = {
        f"{prefix}{WORKBOOK.name}",
        f"{prefix}README_업로드안내.md",
        *(f"{prefix}csv/{path.name}" for path in (PACKAGE / "csv").glob("*.csv")),
    }
    assert names == expected
