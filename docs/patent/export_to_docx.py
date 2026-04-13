"""Export PATENT_STRATEGY_v2_20260406.md to .docx"""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

MD_PATH = Path(__file__).parent / "PATENT_STRATEGY_v2_20260406.md"
OUT_PATH = Path(__file__).parent / "SERS-AI_특허포트폴리오_전략서_v2.1.docx"

doc = Document()

# -- styles --
style = doc.styles['Normal']
font = style.font
font.name = '맑은 고딕'
font.size = Pt(10)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')

for level in range(1, 5):
    hs = doc.styles[f'Heading {level}']
    hs.font.name = '맑은 고딕'
    hs.element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    hs.font.color.rgb = RGBColor(0x1a, 0x1a, 0x2e)

# page margins
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)


def add_table_from_rows(headers, rows):
    """Add a formatted table."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # header
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
        from docx.oxml import OxmlElement
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), 'E8EAF6')
        shading.set(qn('w:val'), 'clear')
        cell._tc.get_or_add_tcPr().append(shading)
    # data
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = val
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph()  # spacing


def parse_table(lines, start_idx):
    """Parse markdown table starting at start_idx, return (headers, rows, end_idx)."""
    # header line
    header_line = lines[start_idx]
    headers = [c.strip() for c in header_line.strip('|').split('|')]
    # separator line
    sep_idx = start_idx + 1
    # data lines
    rows = []
    idx = sep_idx + 1
    while idx < len(lines):
        line = lines[idx].strip()
        if not line.startswith('|'):
            break
        cols = [c.strip() for c in line.strip('|').split('|')]
        rows.append(cols)
        idx += 1
    return headers, rows, idx


def clean_bold_italic(text):
    """Remove markdown bold/italic markers."""
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    text = re.sub(r'~~(.+?)~~', r'\1', text)
    text = re.sub(r'`(.+?)`', r'\1', text)
    return text


def add_rich_paragraph(text, style_name='Normal', bold=False):
    """Add paragraph with inline bold support."""
    p = doc.add_paragraph(style=style_name)
    # split by **...**
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.bold = True
        else:
            # handle ~~strikethrough~~
            sub_parts = re.split(r'(~~.*?~~)', part)
            for sp in sub_parts:
                if sp.startswith('~~') and sp.endswith('~~'):
                    run = p.add_run(sp[2:-2])
                    run.font.strike = True
                else:
                    # handle `code`
                    code_parts = re.split(r'(`.*?`)', sp)
                    for cp in code_parts:
                        if cp.startswith('`') and cp.endswith('`'):
                            run = p.add_run(cp[1:-1])
                            run.font.name = 'Consolas'
                            run.font.size = Pt(9)
                            run.font.color.rgb = RGBColor(0x6a, 0x1b, 0x9a)
                        else:
                            run = p.add_run(cp)
    if bold:
        for run in p.runs:
            run.bold = True
    return p


# -- read and parse --
content = MD_PATH.read_text(encoding='utf-8')
lines = content.split('\n')

i = 0
in_code_block = False
code_lines = []

while i < len(lines):
    line = lines[i]

    # code block
    if line.strip().startswith('```'):
        if in_code_block:
            # end code block
            code_text = '\n'.join(code_lines)
            p = doc.add_paragraph(style='Normal')
            run = p.add_run(code_text)
            run.font.name = 'Consolas'
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            code_lines = []
            in_code_block = False
            i += 1
            continue
        else:
            in_code_block = True
            code_lines = []
            i += 1
            continue

    if in_code_block:
        code_lines.append(line)
        i += 1
        continue

    stripped = line.strip()

    # skip empty lines
    if not stripped:
        i += 1
        continue

    # horizontal rule
    if stripped == '---':
        # skip, we use heading spacing instead
        i += 1
        continue

    # headings
    if stripped.startswith('# ') and not stripped.startswith('## '):
        doc.add_heading(clean_bold_italic(stripped[2:]), level=1)
        i += 1
        continue
    if stripped.startswith('## '):
        doc.add_heading(clean_bold_italic(stripped[3:]), level=2)
        i += 1
        continue
    if stripped.startswith('### '):
        doc.add_heading(clean_bold_italic(stripped[4:]), level=3)
        i += 1
        continue
    if stripped.startswith('#### '):
        doc.add_heading(clean_bold_italic(stripped[5:]), level=4)
        i += 1
        continue

    # blockquote (claims)
    if stripped.startswith('>'):
        quote_lines = []
        while i < len(lines) and lines[i].strip().startswith('>'):
            qt = lines[i].strip().lstrip('>').strip()
            quote_lines.append(qt)
            i += 1
        quote_text = '\n'.join(quote_lines)
        p = doc.add_paragraph(style='Normal')
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.right_indent = Cm(0.5)
        # handle bold in blockquote
        parts = re.split(r'(\*\*.*?\*\*)', quote_text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
                run.font.size = Pt(9.5)
                run.font.italic = True
            else:
                run = p.add_run(part)
                run.font.size = Pt(9.5)
                run.font.italic = True
        continue

    # table
    if stripped.startswith('|') and i + 1 < len(lines) and '---' in lines[i + 1]:
        headers, rows, end_idx = parse_table(lines, i)
        add_table_from_rows(headers, rows)
        i = end_idx
        continue

    # numbered list
    m = re.match(r'^(\d+)\.\s+(.*)', stripped)
    if m:
        add_rich_paragraph(m.group(2), style_name='List Number')
        i += 1
        continue

    # bullet list
    if stripped.startswith('- [ ] ') or stripped.startswith('- [x] '):
        checkbox = '\u2610 ' if '[ ]' in stripped[:6] else '\u2611 '
        text = stripped[6:]
        add_rich_paragraph(checkbox + text, style_name='List Bullet')
        i += 1
        continue

    if stripped.startswith('- '):
        add_rich_paragraph(stripped[2:], style_name='List Bullet')
        i += 1
        continue

    # regular paragraph
    add_rich_paragraph(stripped)
    i += 1

# save
doc.save(str(OUT_PATH))
print(f"Saved: {OUT_PATH}")
