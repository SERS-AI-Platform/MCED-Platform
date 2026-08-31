from __future__ import annotations

import csv
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from ._compat import assert_never
from .clinical_types import (
    ClinicalFormat,
    ClinicalRow,
    ClinicalRowLocator,
    ClinicalSheetContract,
    ClinicalSource,
)
from .clinical_values import ExternalCell, column_names, observation, raw_text


@dataclass(frozen=True, slots=True)
class MissingSubjectKeyError(ValueError):
    source_path: Path
    row_locator: ClinicalRowLocator

    def __str__(self) -> str:
        return f"clinical row has neither patient_id nor solum_label: {self.row_locator}"


@dataclass(frozen=True, slots=True)
class MalformedClinicalSourceError(ValueError):
    source_path: Path

    def __str__(self) -> str:
        return f"clinical source is malformed: {self.source_path.name}"


@dataclass(frozen=True, slots=True)
class MissingClinicalSheetError(ValueError):
    source_path: Path
    sheet_name: str

    def __str__(self) -> str:
        return f"clinical source is missing a configured sheet: {self.source_path.name}"


@dataclass(frozen=True, slots=True)
class MissingClinicalHeaderError(ValueError):
    source_path: Path
    sheet_name: str | None

    def __str__(self) -> str:
        return f"clinical source is missing a configured subject header: {self.source_path.name}"


@dataclass(frozen=True, slots=True)
class ConflictingClinicalIdentityAliasError(ValueError):
    source_path: Path
    sheet_name: str | None

    def __str__(self) -> str:
        return f"clinical source has a conflicting identity alias: {self.source_path.name}"


@dataclass(frozen=True, slots=True)
class _ParsedClinicalRow:
    row: ClinicalRow
    alias_key: str | None


def _clinical_row(
    source: ClinicalSource,
    contract: ClinicalSheetContract,
    locator: ClinicalRowLocator,
    names: tuple[str, ...],
    values: Sequence[ExternalCell],
    identity_aliases: Mapping[str, str],
) -> _ParsedClinicalRow | None:
    populated: list[tuple[str, str]] = []
    for name, value in zip(names, values, strict=False):
        text = raw_text(value)
        if text != "":
            populated.append((name, text))
    if not populated:
        return None
    value_by_name = dict(populated)
    patient_id_field = next(
        (field for field in contract.patient_id_fields if field in value_by_name),
        None,
    )
    patient_id = (
        None if patient_id_field is None else value_by_name.get(patient_id_field)
    )
    solum_label_field = next(
        (
            name
            for name in names
            if name.casefold().replace(" ", "").replace("_", "") == "solumlabel"
        ),
        None,
    )
    solum_label = (
        None if solum_label_field is None else value_by_name.get(solum_label_field)
    )
    if patient_id is None and solum_label is None:
        raise MissingSubjectKeyError(source.path, locator)
    canonical_patient_id = (
        None if patient_id is None else identity_aliases.get(patient_id, patient_id)
    )
    observations = tuple(observation(name, value) for name, value in populated)
    occurred_at = next(
        (
            observation.date_value
            for observation in observations
            if observation.date_value is not None
        ),
        None,
    )
    return _ParsedClinicalRow(
        ClinicalRow(
            locator,
            canonical_patient_id,
            solum_label,
            occurred_at,
            observations,
            contract.event_type,
            patient_id_field,
        ),
        (
            value_by_name.get(contract.subject_key_alias_field)
            if contract.subject_key_alias_field is not None
            else None
        ),
    )


def _default_contract(source: ClinicalSource) -> ClinicalSheetContract:
    return ClinicalSheetContract(
        source.sheet_name,
        source.header_row,
        source.patient_id_fields,
        source.event_type,
    )


def _require_identity_header(
    source: ClinicalSource,
    contract: ClinicalSheetContract,
    headers: tuple[str, ...],
) -> None:
    has_patient_id = any(field in headers for field in contract.patient_id_fields)
    has_solum_label = any(
        header.casefold().replace(" ", "").replace("_", "") == "solumlabel"
        for header in headers
    )
    if not has_patient_id and not has_solum_label:
        raise MissingClinicalHeaderError(source.path, contract.sheet_name)


def _csv_rows(source: ClinicalSource) -> tuple[ClinicalRow, ...]:
    contract = _default_contract(source)
    rows: list[ClinicalRow] = []
    with source.path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream)
        headers = column_names(next(reader, ()))
        _require_identity_header(source, contract, headers)
        for line_number, values in enumerate(reader, start=2):
            parsed = _clinical_row(
                source,
                contract,
                ClinicalRowLocator(f"row:{line_number}"),
                headers,
                values,
                {},
            )
            if parsed is not None:
                rows.append(parsed.row)
    return tuple(rows)


def _excel_rows(source: ClinicalSource) -> tuple[ClinicalRow, ...]:
    try:
        workbook = load_workbook(source.path, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException) as error:
        raise MalformedClinicalSourceError(source.path) from error
    try:
        rows: list[ClinicalRow] = []
        identity_aliases: dict[str, str] = {}
        contracts = source.sheet_contracts or (_default_contract(source),)
        for contract in contracts:
            try:
                worksheet = (
                    workbook[contract.sheet_name]
                    if contract.sheet_name is not None
                    else workbook.active
                )
            except KeyError as error:
                raise MissingClinicalSheetError(
                    source.path,
                    contract.sheet_name or "",
                ) from error
            headers = column_names(
                next(
                    worksheet.iter_rows(
                        min_row=contract.header_row,
                        max_row=contract.header_row,
                        values_only=True,
                    ),
                    (),
                )
            )
            _require_identity_header(source, contract, headers)
            for row_number, values in enumerate(
                worksheet.iter_rows(
                    min_row=contract.data_start_row or contract.header_row + 1,
                    values_only=True,
                ),
                start=contract.data_start_row or contract.header_row + 1,
            ):
                parsed = _clinical_row(
                    source,
                    contract,
                    ClinicalRowLocator(f"{worksheet.title}!row:{row_number}"),
                    headers,
                    values,
                    identity_aliases,
                )
                if parsed is None:
                    continue
                if parsed.alias_key is not None:
                    existing = identity_aliases.get(parsed.alias_key)
                    if existing is not None and existing != parsed.row.patient_id:
                        raise ConflictingClinicalIdentityAliasError(
                            source.path,
                            contract.sheet_name,
                        )
                    if parsed.row.patient_id is not None:
                        identity_aliases[parsed.alias_key] = parsed.row.patient_id
                rows.append(parsed.row)
        return tuple(rows)
    finally:
        workbook.close()


def load_clinical_rows(source: ClinicalSource) -> tuple[ClinicalRow, ...]:
    match source.source_format:
        case ClinicalFormat.CSV:
            return _csv_rows(source)
        case ClinicalFormat.EXCEL:
            return _excel_rows(source)
        case unreachable:
            assert_never(unreachable)
