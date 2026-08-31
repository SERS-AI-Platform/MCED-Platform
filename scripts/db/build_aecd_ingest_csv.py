#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.10"
# dependencies = [
#     "openpyxl>=3.1",
#     "typer>=0.12",
#     "typing-extensions>=4.12",
# ]
# ///

# ─── How to run ───
# 1. Install uv (if not installed):
#      curl -LsSf https://astral.sh/uv/install.sh | sh
# 2. Review defaults and options:
#      uv run scripts/db/build_aecd_ingest_csv.py --help
# 3. Create the CSV:
#      uv run scripts/db/build_aecd_ingest_csv.py
# ──────────────────

# pyright: reportUnnecessaryComparison=false
from __future__ import annotations

import csv
import re
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Annotated, Final, TypeAlias

import typer
from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
from typing_extensions import assert_never, override

ExcelValue: TypeAlias = (
    str
    | int
    | float
    | Decimal
    | bool
    | date
    | datetime
    | time
    | timedelta
    | None
    | ArrayFormula
    | DataTableFormula
    | CellRichText
)

DEFAULT_WORKBOOK: Final = Path("/mnt/c/Users/user/Downloads/전체환자_임상정보_정규화 v5.xlsx")
DEFAULT_MAPPING_ROOT: Final = Path("/home/user/SERS-AI/data/mapping")
DEFAULT_OUTPUT_DIR: Final = Path("/home/user/SERS-AI/data/processed/aecd_platform_ingest")
DEFAULT_SHEET: Final = "table"
DEFAULT_SITE_CODE: Final = "smcxd07"

REQUIRED_COLUMNS: Final = ("group", "solum_label", "patient_code")


@dataclass(frozen=True, slots=True)
class IngestCsvError(Exception):
    message: str

    @override
    def __str__(self) -> str:
        return self.message


@dataclass(frozen=True, slots=True)
class BuildConfig:
    workbook_path: Path
    mapping_root: Path
    sheet_name: str
    site_code: str


@dataclass(frozen=True, slots=True)
class MappingSample:
    source_mapping: str
    folder_label: str


@dataclass(frozen=True, slots=True)
class CsvTable:
    headers: tuple[str, ...]
    rows: tuple[tuple[str, ...], ...]


def normalize_label(value: str) -> str:
    return re.sub(r"[_\s]+", "", value).casefold()


def excel_text(value: ExcelValue) -> str:
    match value:
        case None:
            return ""
        case datetime() as parsed:
            return parsed.date().isoformat()
        case date() as parsed:
            return parsed.isoformat()
        case time() as parsed:
            return parsed.isoformat()
        case timedelta() as parsed:
            return str(parsed)
        case bool() as parsed:
            return "true" if parsed else "false"
        case str() | CellRichText() as parsed:
            return str(parsed).strip()
        case int() | float() | Decimal() as parsed:
            return str(parsed)
        case ArrayFormula() | DataTableFormula():
            raise IngestCsvError("formula cell has no cached scalar value")
        case unreachable:
            assert_never(unreachable)


def scan_mapping(mapping_root: Path) -> tuple[MappingSample, ...]:
    if not mapping_root.is_dir():
        raise IngestCsvError(f"mapping root not found: {mapping_root}")
    samples = tuple(
        MappingSample(mapping_dir.name, sample_dir.name)
        for mapping_dir in sorted(mapping_root.glob("*_mapping"))
        for sample_dir in sorted(mapping_dir.iterdir())
        if sample_dir.is_dir()
    )
    if not samples:
        raise IngestCsvError(f"no sample folders found under: {mapping_root}")
    normalized = [normalize_label(sample.folder_label) for sample in samples]
    duplicates = sorted(label for label, count in Counter(normalized).items() if count > 1)
    if duplicates:
        raise IngestCsvError(f"duplicate mapping labels after normalization: {duplicates}")
    return samples


def read_source_table(workbook_path: Path, sheet_name: str) -> CsvTable:
    if not workbook_path.is_file():
        raise IngestCsvError(f"workbook not found: {workbook_path}")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise IngestCsvError(f"sheet not found: {sheet_name}")
        worksheet = workbook[sheet_name]
        values = worksheet.iter_rows(values_only=True)
        header_values = next(values, None)
        if header_values is None:
            raise IngestCsvError("workbook sheet is empty")
        raw_headers = tuple(excel_text(value) for value in header_values)
        selected_positions = tuple(
            position for position, header in enumerate(raw_headers) if header
        )
        headers = tuple(raw_headers[position] for position in selected_positions)
        duplicates = tuple(header for header, count in Counter(headers).items() if count > 1)
        if duplicates:
            raise IngestCsvError(f"duplicate workbook headers: {duplicates}")
        missing = tuple(column for column in REQUIRED_COLUMNS if column not in headers)
        if missing:
            raise IngestCsvError(f"required workbook columns missing: {missing}")
        label_position = raw_headers.index("solum_label")
        rows = tuple(
            tuple(
                excel_text(row[position] if position < len(row) else None)
                for position in selected_positions
            )
            for row in values
            if excel_text(row[label_position] if label_position < len(row) else None)
        )
        return CsvTable(headers, rows)
    finally:
        workbook.close()


def build_rows(config: BuildConfig) -> CsvTable:
    mappings = scan_mapping(config.mapping_root)
    source_table = read_source_table(config.workbook_path, config.sheet_name)
    positions = {header: position for position, header in enumerate(source_table.headers)}
    label_position = positions["solum_label"]
    patient_position = positions["patient_code"]
    source_by_label: dict[str, tuple[str, ...]] = {}
    duplicate_source_labels: list[str] = []
    for row in source_table.rows:
        label = row[label_position]
        key = normalize_label(label)
        if key in source_by_label:
            duplicate_source_labels.append(label)
        source_by_label[key] = row
    if duplicate_source_labels:
        raise IngestCsvError(
            f"duplicate workbook solum_label values: {sorted(duplicate_source_labels)}"
        )

    unmatched = tuple(
        sample.folder_label
        for sample in mappings
        if normalize_label(sample.folder_label) not in source_by_label
    )
    if unmatched:
        raise IngestCsvError(f"mapping labels not found in workbook: {unmatched}")

    rows = tuple(
        (sample.source_mapping, config.site_code, *source)
        for sample in mappings
        for source in (source_by_label[normalize_label(sample.folder_label)],)
    )
    if len({row[patient_position + 2] for row in rows}) != len(rows):
        raise IngestCsvError("patient_code is not unique in selected rows")
    if len({normalize_label(row[label_position + 2]) for row in rows}) != len(rows):
        raise IngestCsvError("solum_label is not unique in selected rows")
    return CsvTable(("source_mapping", "site_code", *source_table.headers), rows)


def write_rows(output_path: Path, table: CsvTable) -> None:
    if output_path.exists():
        raise IngestCsvError(f"output already exists: {output_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8-sig", newline="") as target:
        writer = csv.writer(target)
        writer.writerow(table.headers)
        writer.writerows(table.rows)


def resolve_output_path(output_file: Path | None, site_code: str) -> Path:
    if output_file is not None:
        return output_file
    return DEFAULT_OUTPUT_DIR / f"{site_code}_clinical_master.csv"


def main(
    input_file: Annotated[Path, typer.Option("--input-file")] = DEFAULT_WORKBOOK,
    mapping_root: Annotated[Path, typer.Option("--mapping-root")] = DEFAULT_MAPPING_ROOT,
    output_file: Annotated[Path | None, typer.Option("--output-file")] = None,
    site_code: Annotated[str, typer.Option("--site-code")] = DEFAULT_SITE_CODE,
    sheet_name: Annotated[str, typer.Option("--sheet-name")] = DEFAULT_SHEET,
) -> None:
    config = BuildConfig(
        workbook_path=input_file,
        mapping_root=mapping_root,
        sheet_name=sheet_name,
        site_code=site_code,
    )
    table = build_rows(config)
    resolved_output = resolve_output_path(output_file, site_code)
    write_rows(resolved_output, table)
    group_position = table.headers.index("group")
    group_counts = dict(Counter(row[group_position] for row in table.rows))
    typer.echo(f"created {resolved_output} with {len(table.rows)} rows; groups={group_counts}")


if __name__ == "__main__":
    typer.run(main)
