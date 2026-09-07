"""임상표 기준 제외 대상을 Lot_Powder DAY 시트에 빨간색으로 표시한다.

  python scripts/db/lot_powder_exclusion/mark_red.py

원본은 건드리지 않고 --out 에 새 파일을 쓴다. DAY 시트 데이터 범위의 기존
빨간색/굵게 표시는 전부 지우고 새로 칠한다. 원본에는 두 종류가 섞여 있으니
(셀 단위 서식 29개, 셀 안 번호만 빨간 rich text 35개) 둘 다 지워야 한다 --
rich text 쪽은 cell.font 로 보이지 않아 처음에 놓쳤던 부분이다.

셀에 번호가 여러 개일 때는 해당 번호만 칠한다. 셀 전체가 제외 대상이면 셀
서식으로, 일부만이면 rich text 로 처리한다.
"""
import argparse
from collections import Counter, defaultdict

import lot_map
import openpyxl
from exclusion_rule import basis_text, evaluate, load, norm, parse_date
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Font
from paths import CLINICAL, LOT, OUT

RED = 'FFFF0000'


def clear_existing(wb):
    """DAY 시트 데이터 범위의 기존 표시를 전부 해제."""
    n_font = n_rich = 0
    for name in lot_map.DAYS:
        ws = wb[name]
        for row in ws.iter_rows(min_row=lot_map.DATA_MIN_ROW, max_row=lot_map.DATA_MAX_ROW,
                                min_col=lot_map.DATA_MIN_COL, max_col=lot_map.DATA_MAX_COL):
            for cell in row:
                if isinstance(cell.value, CellRichText):
                    cell.value = ''.join(str(getattr(b, 'text', b)) for b in cell.value)
                    n_rich += 1
                f = cell.font
                try:
                    rgb = f.color.rgb if f.color else None
                except Exception:
                    rgb = None
                if (isinstance(rgb, str) and rgb.upper().endswith('FF0000')) or f.bold:
                    cell.font = Font(name=f.name, size=f.sz, bold=False,
                                     italic=f.i, underline=f.u, color=None)
                    n_font += 1
    return n_font, n_rich


def mark(wb, hits, selected):
    """제외 대상 번호만 빨간색으로. [(sheet, coord, number, label)] 반환."""
    by_cell = defaultdict(list)
    for sheet, coord, _group, num, label in hits:
        by_cell[(sheet, coord)].append((num, label))

    marked, n_whole, n_partial = [], 0, 0
    for (sheet, coord), tokens in by_cell.items():
        hit = [t for t in tokens if t[1] in selected]
        if not hit:
            continue
        cell = wb[sheet][coord]
        base = cell.font
        if len(hit) == len(tokens):
            cell.font = Font(name=base.name, size=base.sz, bold=True, color=RED)
            n_whole += 1
        else:
            import re
            hit_nums = {str(n) for n, _ in hit}
            blocks = []
            for part in re.split(r'([,\s]+)', str(cell.value)):
                if part.strip() in hit_nums:
                    blocks.append(TextBlock(
                        InlineFont(rFont=base.name, sz=base.sz, b=True, color=RED), part))
                elif part:
                    blocks.append(part)
            cell.value = CellRichText(blocks)
            n_partial += 1
        marked += [(sheet, coord, n, label) for n, label in hit]
    return marked, n_whole, n_partial


def write_sheets(wb, res, ix, rows, selected, marked, clinical_path):
    """근거 시트 두 개를 추가한다. 기존 제외사유_목록/요약은 건드리지 않는다."""
    prefix2group = {norm(v): k for k, v in lot_map.GROUP2PREFIX.items()}
    cancer_type = {norm(v[ix['solum_label']]): v[ix['cancer_type']] for v, _, _ in rows}
    timing = {norm(v[ix['solum_label']]): v[ix['primary_sample_timing_interpretation']]
              for v, _, _ in rows}
    dates = {norm(v[ix['solum_label']]): tuple(
        parse_date(v[ix[k]]) for k in ('collection_date', 'surgery_date', 'chemo_start_date'))
        for v, _, _ in rows}

    ws = wb.create_sheet('제외사유_목록_v8')
    ws.append(['No.', 'Lot_Sheet', 'Lot_Cell', 'Lot_Group', 'Lot_번호', 'solum_label',
               'cancer_type', 'sample_timing', 'collection_date', 'surgery_date',
               'chemo_start_date', '요로감염', '타암이력', '치료후채취', '다중암',
               'Drop', '노란표시', '제외근거'])
    for i, (sheet, coord, num, label) in enumerate(sorted(marked), 1):
        r = res[label]
        ws.append([i, sheet, coord,
                   prefix2group.get(label.rsplit('_', 1)[0], label.rsplit('_', 1)[0]),
                   num, label, cancer_type.get(label), timing.get(label),
                   *dates.get(label, (None, None, None)),
                   'Y' if r['uti'] else '', 'Y' if r['others'] else '',
                   'Y' if r['post'] else '', 'Y' if r['multi'] else '',
                   'Y' if r['drop'] else '', 'Y' if r['yellow'] else '', basis_text(r)])

    ws2 = wb.create_sheet('제외사유_요약_v8')
    ws2.append(['항목', '건수', '설명'])
    ws2.append(['요로감염', sum(1 for label in selected if res[label]['uti']),
                'ua_nitrite = Positive'])
    ws2.append(['  - Nitrite + LE 동시 Positive',
                sum(1 for label in selected if res[label]['both']), '위 중 LE 도 Positive 인 건'])
    ws2.append(['  - Nitrite 단독 Positive',
                sum(1 for label in selected if res[label]['uti'] and not res[label]['both']), ''])
    ws2.append(['타암 이력', sum(1 for label in selected if res[label]['others']),
                'past_history_1~9 에 현재 암종과 다른 암/종양'])
    ws2.append(['치료 후 채취', sum(1 for label in selected if res[label]['post']),
                'collection_date 가 surgery/chemo 시작일보다 뒤'
                ' (pre_treatment_specimen 태그가 있으면 제외 안 함)'])
    ws2.append(['다중암 환자', sum(1 for label in selected if res[label]['multi']),
                '19022041 (BLC_247, PRO_60) / 19276017 (CRC_186, PAN_78)'])
    ws2.append(["group='Drop'", sum(1 for label in selected if res[label]['drop']),
                '임상표 table 시트에서 초록색으로 칠해진 행'])
    ws2.append(['임상표 노란 표시 (사유 미상)',
                sum(1 for label in selected if res[label]['yellow'] and not (
                    res[label]['uti'] or res[label]['others'] or res[label]['post']
                    or res[label]['multi'] or res[label]['drop'])),
                '노란 표시 외에 다른 근거가 없는 건'])
    ws2.append(['최종 제외 라벨 (합집합)', len(selected), '위 근거들의 합집합'])
    ws2.append(['DAY 시트에 표시한 번호', len(marked), 'Lot 셀 안에서 빨간색으로 칠한 번호 수'])
    ws2.append(['소스 임상표', str(clinical_path), 'table 시트'])
    ws2.append(['판정 제외', 'KPAN', '재료연 검체 -- 임상표에 대응 라벨이 없어 표시 대상 아님'])
    ws2.append(['그룹별', ', '.join(f'{k} {v}' for k, v in sorted(
        Counter(label.rsplit('_', 1)[0] for label in selected).items())), ''])


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--clinical', default=CLINICAL, help='임상정보 정규화 워크북')
    ap.add_argument('--lot', default=LOT, help='Lot_Powder 워크북 (원본, 수정하지 않음)')
    ap.add_argument('--out', default=OUT, help='저장할 새 워크북')
    a = ap.parse_args()

    ix, rows = load(a.clinical)
    res = evaluate(ix, rows)
    hits = lot_map.scan(a.lot)
    lot_labels = {h[4] for h in hits}
    missing = sorted(lot_labels - set(res))
    if missing:
        print(f'경고: 임상표에 없는 Lot 라벨 {len(missing)}개 -- {missing[:10]}')
    selected = {label for label in lot_labels if res.get(label, {}).get('excl')}
    print(f'판정 대상 Lot 라벨 {len(lot_labels)} | 제외 대상 {len(selected)}')

    wb = openpyxl.load_workbook(a.lot, rich_text=True)
    n_font, n_rich = clear_existing(wb)
    print(f'기존 표시 해제: 셀 단위 {n_font} / rich text {n_rich}')
    marked, n_whole, n_partial = mark(wb, hits, selected)
    print(f'셀 전체 표시 {n_whole} / 일부 번호만 {n_partial} / 표시한 번호 {len(marked)}')
    write_sheets(wb, res, ix, rows, selected, marked, a.clinical)
    wb.save(a.out)
    print('저장:', a.out)


if __name__ == '__main__':
    main()
