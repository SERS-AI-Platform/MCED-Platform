"""mark_red.py 가 저장한 파일을 다시 열어 표시가 정확한지 검사한다.

  python scripts/db/lot_powder_exclusion/verify_marks.py

세 가지를 본다:
  1) 빨간 번호 집합이 규칙이 고른 집합과 정확히 같은가 (누락/과다 0)
  2) 원본 대비 텍스트와 수식이 하나도 안 바뀌었는가 (서식만 바뀌어야 한다)
  3) KPAN 열에 빨간 표시가 남아 있지 않은가
"""
import argparse

import lot_map
import openpyxl
from exclusion_rule import evaluate, load
from openpyxl.cell.rich_text import CellRichText
from openpyxl.worksheet.formula import ArrayFormula
from paths import CLINICAL, LOT, OUT


def is_red(font):
    try:
        return font is not None and font.color is not None \
            and str(font.color.rgb).upper().endswith('FF0000')
    except Exception:
        return False


def red_numbers(wb, only_group=None):
    out = set()
    for name in lot_map.DAYS:
        ws = wb[name]
        header = {c.column: c.value for c in ws[2] if c.value}
        for row in ws.iter_rows(min_row=lot_map.DATA_MIN_ROW, max_row=lot_map.DATA_MAX_ROW):
            for cell in row:
                group = header.get(cell.column)
                if only_group is None:
                    if group not in lot_map.GROUP2PREFIX:
                        continue
                elif group != only_group:
                    continue
                v = cell.value
                if isinstance(v, CellRichText):
                    for b in v:
                        if is_red(getattr(b, 'font', None)):
                            for t in str(getattr(b, 'text', b)).split(','):
                                if t.strip().isdigit():
                                    out.add((name, cell.coordinate, int(t.strip())))
                elif is_red(cell.font):
                    for t in str(v).split(','):
                        if t.strip().isdigit():
                            out.add((name, cell.coordinate, int(t.strip())))
    return out


def as_text(v):
    if isinstance(v, CellRichText):
        return ''.join(str(getattr(t, 'text', t)) for t in v)
    if isinstance(v, ArrayFormula):
        return 'AF:' + str(v.text)
    return '' if v is None else str(v)


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--clinical', default=CLINICAL)
    ap.add_argument('--lot', default=LOT)
    ap.add_argument('--out', default=OUT)
    a = ap.parse_args()

    ix, rows = load(a.clinical)
    res = evaluate(ix, rows)
    hits = lot_map.scan(a.lot)
    want = {(s, c, n) for s, c, _g, n, label in hits if res.get(label, {}).get('excl')}

    src = openpyxl.load_workbook(a.lot, rich_text=True)
    dst = openpyxl.load_workbook(a.out, rich_text=True)

    got = red_numbers(dst)
    print(f'1) 의도한 번호 {len(want)} | 실제 빨강 {len(got)}'
          f' | 누락 {len(want - got)} | 과다 {len(got - want)}')
    for x in sorted(want - got)[:10]:
        print('   누락', x)
    for x in sorted(got - want)[:10]:
        print('   과다', x)

    diff = 0
    for name in src.sheetnames:
        x, y = src[name], dst[name]
        for r in range(1, max(x.max_row, y.max_row) + 1):
            for c in range(1, max(x.max_column, y.max_column) + 1):
                if as_text(x.cell(r, c).value) != as_text(y.cell(r, c).value):
                    diff += 1
    print(f'2) 원본 대비 텍스트/수식 변경 셀 {diff}')

    print(f'3) KPAN 열에 남은 빨간 표시 {len(red_numbers(dst, only_group="KPAN"))}')


if __name__ == '__main__':
    main()
