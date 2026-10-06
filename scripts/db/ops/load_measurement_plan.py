"""Load the measurement plan workbook into ops.measurement_plan / ops.clinical_exclusions (idempotent).

The workbook ("검체 측정 Lot_Powder_…xlsx") has one sheet per measuring day (DAY 1..n): column A = Lot (carried
down), column B = Paper '#n', row 2 = group headers, cells = sample numbers ("57,58"). Cell font colours are
not used -- red marks are partial rich text; the exclusion sheet ("제외사유_목록_…") is the source instead.
A file is identified by its MD5: loading the same file again only re-marks it current.
--previous loads an older plan's cells only (no exclusion sheet, current plan untouched); v_plan_progress adds its
cells that the current plan lacks as plan_scope '이전 계획', so the dashboard shows overall progress.

    source scripts/db/pghost.sh
    python scripts/db/ops/load_measurement_plan.py --file "<xlsx>" [--dry-run] [--previous]
"""

from __future__ import annotations

import argparse
import hashlib
import re
import sys
from pathlib import Path

import openpyxl
import psycopg2
from psycopg2.extras import execute_values

GROUP_ALIAS = {"CPAN": "PAN"}          # 계획표 열 이름 -> master.samples 라벨 접두어
EXCL_FLAGS = ("요로감염", "타암이력", "치료후채취", "다중암", "Drop", "노란표시")


def label(prefix: str, n: str | int) -> str:
    return f"{GROUP_ALIAS.get(prefix, prefix)}_{n}"


def parse_plan(wb) -> list[dict]:
    rows = []
    for ws in wb.worksheets:
        if not ws.title.startswith("DAY"):
            continue
        header = [c.value for c in ws[2]]
        lot = None
        for r in ws.iter_rows(min_row=3):
            if r[0].value:
                lot = str(r[0].value).strip()
            paper = r[1].value
            if not (paper and str(paper).startswith("#")):
                continue
            for j, grp in enumerate(header):
                if j < 2 or grp in (None, "Total") or r[j].value is None:
                    continue
                for n in re.findall(r"\d+", str(r[j].value)):
                    rows.append({"label": label(grp, n), "day": ws.title, "lot": lot, "paper": str(paper),
                                 "group": grp, "cell": r[j].coordinate})
    return rows


def parse_exclusions(wb) -> tuple[list[dict], str]:
    ws = next(s for s in wb.worksheets if s.title.startswith("제외사유_목록"))
    criteria = ws.title.rsplit("_", 1)[-1]                  # 예: 'v8'
    head = [c.value for c in ws[1]]
    ix = {h: i for i, h in enumerate(head)}
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[ix["solum_label"]]:
            continue
        lab = str(r[ix["solum_label"]]).replace("H.D._", "H. D._")   # DB 라벨은 'H. D._n'
        out.append({"label": lab, "reason": r[ix["제외근거"]],
                    **{f: str(r[ix[f]] or "").strip().upper() == "Y" for f in EXCL_FLAGS}})
    return out, criteria


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--file", type=Path, required=True)
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--previous", action="store_true",
                    help="이전 계획표: 계획 칸만 적재하고 현재 계획은 바꾸지 않음 (제외 목록은 현재 계획표 것을 씀)")
    args = ap.parse_args()
    md5 = hashlib.md5(args.file.read_bytes()).hexdigest()
    wb = openpyxl.load_workbook(args.file, data_only=True)
    plan = parse_plan(wb)
    excl, criteria = ([], "-") if args.previous else parse_exclusions(wb)

    labels = [p["label"] for p in plan]
    dup = {x for x in labels if labels.count(x) > 1}
    if dup:
        sys.exit(f"계획표에 같은 검체가 두 번 이상: {sorted(dup)[:10]}")

    with psycopg2.connect(dbname="aecd_platform") as conn, conn.cursor() as cur:
        cur.execute("SELECT solum_label, sample_id FROM master.samples")
        sid = dict(cur.fetchall())
        missing = sorted(({p["label"] for p in plan} | {e["label"] for e in excl}) - set(sid))
        if missing:
            sys.exit(f"DB에 없는 라벨 {len(missing)}개: {missing[:10]}")
        print(f"{args.file.name} (md5 {md5[:8]}): 계획 {len(plan)}칸 · 시트 "
              f"{len({p['day'] for p in plan})}개 · Lot {len({p['lot'] for p in plan})}개 · 제외 목록 {len(excl)}건({criteria}), "
              f"그중 계획에 있는 것 {len({e['label'] for e in excl} & set(labels))}", flush=True)
        if args.dry_run:
            print("dry-run: 쓰지 않음")
            return
        if args.previous:
            cur.execute("SELECT 1 FROM ops.plan_versions WHERE file_md5 = %s AND is_current", (md5,))
            if cur.fetchone():
                sys.exit("이 파일은 현재 계획이다 — --previous로 다시 넣으면 제외 목록이 지워짐")

        cur.execute("""INSERT INTO ops.plan_versions (source_file, file_md5) VALUES (%s, %s)
                       ON CONFLICT (file_md5) DO UPDATE SET source_file = EXCLUDED.source_file
                       RETURNING plan_version_id""", (str(args.file), md5))
        vid = cur.fetchone()[0]
        cur.execute("DELETE FROM ops.measurement_plan WHERE plan_version_id = %s", (vid,))
        cur.execute("DELETE FROM ops.clinical_exclusions WHERE plan_version_id = %s", (vid,))
        execute_values(cur, """INSERT INTO ops.measurement_plan
            (plan_version_id, sample_id, solum_label, plan_day, plan_lot, paper_no, plan_group, sheet_cell) VALUES %s""",
            [(vid, sid[p["label"]], p["label"], p["day"], p["lot"], p["paper"], p["group"], p["cell"]) for p in plan])
        execute_values(cur, """INSERT INTO ops.clinical_exclusions
            (plan_version_id, sample_id, solum_label, urinary_infection, other_cancer, post_treatment, multi_cancer,
             drop_flag, yellow_mark, reason, criteria) VALUES %s""",
            [(vid, sid[e["label"]], e["label"], *(e[f] for f in EXCL_FLAGS), e["reason"], criteria) for e in excl])
        if not args.previous:
            cur.execute("UPDATE ops.plan_versions SET is_current = false WHERE is_current AND plan_version_id <> %s", (vid,))
            cur.execute("UPDATE ops.plan_versions SET is_current = true WHERE plan_version_id = %s", (vid,))
        conn.commit()

        cur.execute("""SELECT plan_scope, count(*), count(*) FILTER (WHERE measured), count(*) FILTER (WHERE excluded)
                       FROM ops.v_plan_progress GROUP BY 1 ORDER BY 1 DESC""")
        scopes = {s: (n, m, x) for s, n, m, x in cur.fetchall()}
        for s, (n, m, x) in scopes.items():
            print(f"적재 확인 [{s}]: {n}칸 · BCCP0922 측정 {m} · 제외 대상 {x}", flush=True)
        if not args.previous and scopes.get("현재 계획", (0,))[0] != len(plan):
            sys.exit("v_plan_progress 현재 계획 행 수가 계획 칸 수와 다름 (조인 중복?)")


if __name__ == "__main__":
    main()
