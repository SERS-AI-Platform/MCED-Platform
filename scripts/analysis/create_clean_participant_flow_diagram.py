from __future__ import annotations

from pathlib import Path
from xml.sax.saxutils import escape

import cairosvg
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

W, H = 1800, 1350
PROJECT_ROOT = Path(__file__).resolve().parents[2]
OUT = PROJECT_ROOT / "results" / "clean_cohort_20260605" / "clean_participant_flow_diagram"
FONT = "DejaVu Sans, Arial, Helvetica, sans-serif"

BLACK = "#000000"
DARK = "#222222"
MID = "#555555"
PALE = "#F5F5F5"
WHITE = "#FFFFFF"


def svg_text(x, y, text, size=30, weight=400, fill=DARK, anchor="middle"):
    return (
        f'<text x="{x}" y="{y}" text-anchor="{anchor}" '
        f'font-family="{FONT}" font-size="{size}" font-weight="{weight}" '
        f'fill="{fill}">{escape(text)}</text>'
    )


def svg_multiline(x, y, lines, size=26, weight=400, fill=DARK, anchor="middle", leading=1.3):
    return "\n".join(
        svg_text(x, y + i * size * leading, line, size, weight, fill, anchor)
        for i, line in enumerate(lines)
    )


def svg_rect(x, y, w, h, fill=WHITE, stroke=BLACK, sw=2.4, dash=None):
    dash_attr = f' stroke-dasharray="{dash}"' if dash else ""
    return (
        f'<rect x="{x}" y="{y}" width="{w}" height="{h}" fill="{fill}" '
        f'stroke="{stroke}" stroke-width="{sw}"{dash_attr}/>'
    )


def svg_line(x1, y1, x2, y2, sw=2.4, dash=None, arrow=False):
    dash_attr = f' stroke-dasharray="{dash}"' if dash else ""
    marker = ' marker-end="url(#arrow)"' if arrow else ""
    return (
        f'<line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" '
        f'stroke="{BLACK}" stroke-width="{sw}" fill="none"{dash_attr}{marker}/>'
    )


def svg_path(d, sw=1.8, dash=None):
    dash_attr = f' stroke-dasharray="{dash}"' if dash else ""
    return f'<path d="{d}" stroke="{BLACK}" stroke-width="{sw}" fill="none"{dash_attr}/>'


def main_box(x, y, w, h, title, n_text, details):
    cx = x + w / 2
    items = [svg_rect(x, y, w, h)]
    items.append(svg_text(cx, y + 42, title, 30, 700, BLACK))
    items.append(svg_text(cx, y + 94, n_text, 44, 700, BLACK))
    items.append(svg_multiline(cx, y + 132, details, 21, 400, MID, "middle", 1.32))
    return "\n".join(items)


def exclusion_box(x, y, w, h, title, lines):
    items = [svg_rect(x, y, w, h, WHITE, BLACK, 2.0, "7 7")]
    items.append(svg_text(x + 28, y + 38, title, 24, 700, BLACK, "start"))
    items.append(svg_multiline(x + 28, y + 78, lines, 20, 400, DARK, "start", 1.42))
    return "\n".join(items)


def hospital_strip(x, y, w, h):
    hospitals = [
        ("CBNUH", "n = 638", ["BLC 208 | CRC 285", "CPAN 54 | PRO 91"]),
        ("YPNUH", "n = 372", ["Controls: NOR, DIA, HBP, H.D."]),
        ("SNUH", "n = 170", ["LUN 130 | OVA 40"]),
        ("SSMH", "n = 163", ["LUN 163"]),
        ("IJBPH", "n = 52", ["BRE 29 | OVA 23"]),
    ]
    cell_w = w / len(hospitals)
    items = [
        svg_rect(x, y, w, h),
        svg_rect(x, y, w, 52, PALE, BLACK, 2.0),
        svg_text(x + 26, y + 35, "Final model cohort by hospital", 24, 700, BLACK, "start"),
        svg_text(x + w - 26, y + 35, "N = 1,395 subjects | 6,975 spectra", 23, 700, BLACK, "end"),
    ]
    for i, (code, n_text, details) in enumerate(hospitals):
        cx = x + i * cell_w
        items.append(svg_rect(cx, y + 78, cell_w, h - 78, WHITE, BLACK, 1.4))
        items.append(svg_text(cx + cell_w / 2, y + 106, code, 24, 700, BLACK))
        items.append(svg_text(cx + cell_w / 2, y + 140, n_text, 26, 700, BLACK))
        items.append(svg_multiline(cx + cell_w / 2, y + 168, details, 14, 400, MID, "middle", 1.2))
    return "\n".join(items)


def build_svg():
    top_x = 510
    box_w = 780
    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" viewBox="0 0 {W} {H}">',
        "<defs>",
        '<marker id="arrow" markerWidth="14" markerHeight="14" refX="11" refY="7" orient="auto" markerUnits="strokeWidth">',
        f'<path d="M2,2 L12,7 L2,12 Z" fill="{BLACK}"/>',
        "</marker>",
        "</defs>",
        f'<rect width="100%" height="100%" fill="{WHITE}"/>',
        svg_text(900, 62, "Clean Retrospective Cohort Flow Diagram", 40, 700, BLACK),
        svg_text(
            900,
            101,
            "Selection of eligible subjects for the final SERS model analysis cohort",
            23,
            400,
            MID,
        ),
        main_box(
            top_x,
            125,
            box_w,
            154,
            "Target retrospective clinical cohort",
            "N = 1,570",
            ["Seven target cancer groups and controls after harmonization"],
        ),
        svg_line(900, 279, 900, 358, 3.0, arrow=True),
        exclusion_box(
            1305,
            245,
            430,
            182,
            "Excluded at step 1",
            [
                "Not effective/pre-treatment: N = 66",
                "BLC: 48; CPAN: 10",
                "LUN: 7; CRC: 1",
                "Controls: 0",
            ],
        ),
        svg_path("M1305 336 C1185 330 1095 310 965 308", 1.8, "8 8"),
        main_box(
            top_x,
            378,
            box_w,
            168,
            "Effective pre-treatment specimen or control",
            "N = 1,504",
            ["Cancer: n = 1,104 | Controls: n = 400", "Retained if is_pre_effective = 1"],
        ),
        svg_line(900, 546, 900, 626, 3.0, arrow=True),
        exclusion_box(
            80,
            520,
            450,
            196,
            "Excluded at step 2",
            [
                "Other cancer history: N = 108",
                "Cancer with different cancer history: 80",
                "Controls with prior cancer history: 28",
                "DIA: 7; HBP: 5; H.D.: 16",
            ],
        ),
        svg_path("M530 622 C630 622 724 626 835 626", 1.8, "8 8"),
        main_box(
            top_x,
            650,
            box_w,
            186,
            "Clean clinical cohort",
            "N = 1,396",
            ["Cancer: n = 1,024 | Controls: n = 372", "No other-cancer history"],
        ),
        svg_line(900, 836, 900, 916, 3.0, arrow=True),
        exclusion_box(
            1305,
            790,
            430,
            128,
            "Excluded at step 3",
            [
                "No matched processed spectrum: N = 1",
                "One bladder cancer subject",
            ],
        ),
        svg_path("M1305 852 C1182 854 1096 872 965 878", 1.8, "8 8"),
        main_box(
            top_x,
            940,
            box_w,
            170,
            "Final model analysis cohort",
            "N = 1,395",
            ["Cancer: n = 1,023 | Controls: n = 372", "Processed SERS spectra: n = 6,975"],
        ),
        hospital_strip(80, 1140, 1640, 190),
        "</svg>",
    ]
    return "\n".join(parts)


def ppt_color(hex_color):
    hex_color = hex_color.lstrip("#")
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def set_shape_style(shape, fill=WHITE, line=BLACK, width=1.5, dash=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = ppt_color(fill)
    shape.line.color.rgb = ppt_color(line)
    shape.line.width = Pt(width)
    if dash:
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH


def build_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(10.0)
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    scale_x = prs.slide_width / W
    scale_y = prs.slide_height / H

    def emu_x(px):
        return int(px * scale_x)

    def emu_y(px):
        return int(px * scale_y)

    def add_rect(x, y, w, h, fill=WHITE, sw=1.5, dash=None):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            emu_x(x),
            emu_y(y),
            emu_x(w),
            emu_y(h),
        )
        set_shape_style(shape, fill, BLACK, sw, dash)
        return shape

    def add_text(x, y, w, h, text, size=20, weight=False, color=DARK, align="center"):
        tb = slide.shapes.add_textbox(emu_x(x), emu_y(y), emu_x(w), emu_y(h))
        tf = tb.text_frame
        tf.clear()
        tf.margin_left = 0
        tf.margin_right = 0
        tf.margin_top = 0
        tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = {"center": PP_ALIGN.CENTER, "left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT}[
            align
        ]
        run = p.add_run()
        run.text = text
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.bold = weight
        run.font.color.rgb = ppt_color(color)
        return tb

    def add_multiline(x, y, w, h, lines, size=13, color=MID, align="center"):
        tb = slide.shapes.add_textbox(emu_x(x), emu_y(y), emu_x(w), emu_y(h))
        tf = tb.text_frame
        tf.clear()
        tf.margin_left = 0
        tf.margin_right = 0
        tf.margin_top = 0
        tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {
                "center": PP_ALIGN.CENTER,
                "left": PP_ALIGN.LEFT,
                "right": PP_ALIGN.RIGHT,
            }[align]
            run = p.add_run()
            run.text = line
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.font.color.rgb = ppt_color(color)
        return tb

    def add_line(x1, y1, x2, y2, width=1.5, dash=False, arrow=False):
        conn = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, emu_x(x1), emu_y(y1), emu_x(x2), emu_y(y2)
        )
        conn.line.color.rgb = ppt_color(BLACK)
        conn.line.width = Pt(width)
        if dash:
            conn.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        if arrow:
            conn.line.end_arrowhead = True
        return conn

    def ppt_main_box(x, y, w, h, title, n_text, details):
        add_rect(x, y, w, h, WHITE, 1.8)
        add_text(x + 30, y + 18, w - 60, 38, title, 18, True, BLACK)
        add_text(x + 30, y + 64, w - 60, 48, n_text, 28, True, BLACK)
        add_multiline(x + 36, y + 112, w - 72, h - 118, details, 12.5, MID)

    def ppt_exclusion(x, y, w, h, title, lines):
        add_rect(x, y, w, h, WHITE, 1.2, dash=True)
        add_text(x + 28, y + 12, w - 56, 34, title, 14.5, True, BLACK, "left")
        add_multiline(x + 28, y + 54, w - 56, h - 64, lines, 11.2, DARK, "left")

    add_text(0, 42, W, 42, "Clean Retrospective Cohort Flow Diagram", 24, True, BLACK)
    add_text(
        0,
        82,
        W,
        30,
        "Selection of eligible subjects for the final SERS model analysis cohort",
        13.5,
        False,
        MID,
    )

    top_x = 510
    box_w = 780
    ppt_main_box(
        top_x,
        125,
        box_w,
        154,
        "Target retrospective clinical cohort",
        "N = 1,570",
        ["Seven target cancer groups and controls after harmonization"],
    )
    add_line(900, 279, 900, 358, 2.0, arrow=True)
    ppt_exclusion(
        1305,
        245,
        430,
        182,
        "Excluded at step 1",
        [
            "Not effective/pre-treatment: N = 66",
            "BLC: 48; CPAN: 10",
            "LUN: 7; CRC: 1",
            "Controls: 0",
        ],
    )
    add_line(1305, 336, 965, 308, 1.0, dash=True)
    ppt_main_box(
        top_x,
        378,
        box_w,
        168,
        "Effective pre-treatment specimen or control",
        "N = 1,504",
        ["Cancer: n = 1,104 | Controls: n = 400", "Retained if is_pre_effective = 1"],
    )
    add_line(900, 546, 900, 626, 2.0, arrow=True)
    ppt_exclusion(
        80,
        520,
        450,
        196,
        "Excluded at step 2",
        [
            "Other cancer history: N = 108",
            "Cancer with different cancer history: 80",
            "Controls with prior cancer history: 28",
            "DIA: 7; HBP: 5; H.D.: 16",
        ],
    )
    add_line(530, 622, 835, 626, 1.0, dash=True)
    ppt_main_box(
        top_x,
        650,
        box_w,
        186,
        "Clean clinical cohort",
        "N = 1,396",
        ["Cancer: n = 1,024 | Controls: n = 372", "No other-cancer history"],
    )
    add_line(900, 836, 900, 916, 2.0, arrow=True)
    ppt_exclusion(
        1305,
        790,
        430,
        128,
        "Excluded at step 3",
        ["No matched processed spectrum: N = 1", "One bladder cancer subject"],
    )
    add_line(1305, 852, 965, 878, 1.0, dash=True)
    ppt_main_box(
        top_x,
        940,
        box_w,
        170,
        "Final model analysis cohort",
        "N = 1,395",
        ["Cancer: n = 1,023 | Controls: n = 372", "Processed SERS spectra: n = 6,975"],
    )

    x, y, w, h = 80, 1140, 1640, 190
    add_rect(x, y, w, h, WHITE, 1.5)
    add_rect(x, y, w, 52, PALE, 1.2)
    add_text(x + 26, y + 10, 500, 34, "Final model cohort by hospital", 14.5, True, BLACK, "left")
    add_text(
        x + w - 520, y + 10, 495, 34, "N = 1,395 subjects | 6,975 spectra", 14, True, BLACK, "right"
    )

    hospitals = [
        ("CBNUH", "n = 638", ["BLC 208 | CRC 285", "CPAN 54 | PRO 91"]),
        ("YPNUH", "n = 372", ["Controls: NOR, DIA, HBP, H.D."]),
        ("SNUH", "n = 170", ["LUN 130 | OVA 40"]),
        ("SSMH", "n = 163", ["LUN 163"]),
        ("IJBPH", "n = 52", ["BRE 29 | OVA 23"]),
    ]
    cell_w = w / len(hospitals)
    for i, (code, n_text, details) in enumerate(hospitals):
        cx = x + i * cell_w
        add_rect(cx, y + 78, cell_w, h - 78, WHITE, 1.0)
        add_text(cx, y + 82, cell_w, 32, code, 14, True, BLACK)
        add_text(cx, y + 116, cell_w, 34, n_text, 15, True, BLACK)
        add_multiline(cx + 8, y + 148, cell_w - 16, 44, details, 8.4, MID)

    prs.save(f"{OUT}.pptx")


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    svg = build_svg()
    Path(f"{OUT}.svg").write_text(svg, encoding="utf-8")
    cairosvg.svg2png(
        bytestring=svg.encode("utf-8"), write_to=f"{OUT}.png", output_width=W, output_height=H
    )
    build_pptx()
    print(f"Wrote {OUT}.svg")
    print(f"Wrote {OUT}.png")
    print(f"Wrote {OUT}.pptx")


if __name__ == "__main__":
    main()
