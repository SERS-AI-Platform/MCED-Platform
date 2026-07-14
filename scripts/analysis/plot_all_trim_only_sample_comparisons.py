#!/usr/bin/env python3
"""Create trim-only primary-vs-clinical raw comparison plots for all samples.

This script reads native CSV values directly and applies only the requested
x-range trim. It does not interpolate, average, sort, smooth, baseline-correct,
normalize, reference-correct, blank-correct, or compute centroids.
"""

from __future__ import annotations

import argparse
import math
import re
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from trim_only_raw_utils import (
    DEFAULT_CLINICAL_ROOT,
    DEFAULT_PRIMARY_ROOT,
    collect_control_files,
    date_key_from_path,
    date_root_from_path,
    is_average_file,
    normalize_group,
    parse_sample,
    read_trimmed_xy,
    readable_spectrum_files,
    sample_files,
)

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUT_DIR = (
    PROJECT_ROOT
    / "results"
    / "raw_data_vs_clinical_reference"
    / "figures"
    / "trim_only_all_samples"
)

GROUP_ORDER = [
    "NOR",
    "YNOR",
    "DIA",
    "HBP",
    "H.D.",
    "CPAN",
    "YPAN",
    "PRO",
    "OVA",
    "BRE",
    "BLC",
    "CRC",
    "LUN",
]
GROUP_COLORS = {
    "NOR": "#2f6f4e",
    "YNOR": "#77a65a",
    "DIA": "#b07d2c",
    "HBP": "#8e5ea2",
    "H.D.": "#7c5f45",
    "CPAN": "#d17c4a",
    "YPAN": "#e0a33a",
    "PRO": "#486fb5",
    "OVA": "#be5c8a",
    "BRE": "#c75f55",
    "BLC": "#4f8f9f",
    "CRC": "#c04537",
    "LUN": "#344f7a",
}


@dataclass(frozen=True)
class PaperKey:
    lot: int
    paper: int

    @property
    def label(self) -> str:
        return f"Lot.{self.lot} Paper #{self.paper}"


@dataclass(frozen=True)
class SamplePlan:
    group: str
    sample_id: str
    paper_key: PaperKey
    paper_index: int
    reference_batch: int


@dataclass(frozen=True)
class ReferenceControl:
    requested_batch: int
    source_batch: int
    paths: list[Path]

    @property
    def carried_forward(self) -> bool:
        return self.requested_batch != self.source_batch


BLANK_FILE_RE = re.compile(r"(?i)^Blank\s+(?P<lot>\d+)-(?P<paper>\d+)_(?P<rep>\d+|ave)$")
LOT_RE = re.compile(r"(?i)Lot\.?\s*(\d+)")
PAPER_RE = re.compile(r"#\s*(\d+)")


def safe_label(group: str, sample_id: str, date_key: str) -> str:
    safe_group = group.replace(".", "D").replace(" ", "")
    return f"{safe_group}_{sample_id}_{date_key}"


def sort_key(item: dict) -> tuple[int, str, int, str]:
    group = item["group"]
    group_rank = GROUP_ORDER.index(group) if group in GROUP_ORDER else len(GROUP_ORDER)
    return group_rank, item["date_key"], int(item["sample_id"]), group


def normalize_plan_group(group: object) -> str | None:
    if group is None:
        return None
    text = str(group).strip()
    if not text or text.lower() == "total":
        return None
    return normalize_group(text)


def numeric_ids_from_cell(value: object, ignore_parentheses: bool = True) -> list[str]:
    if value is None:
        return []
    if isinstance(value, float) and value.is_integer():
        return [str(int(value))]
    if isinstance(value, int):
        return [str(value)]
    text = str(value)
    if ignore_parentheses:
        text = re.sub(r"\([^)]*\)", "", text)
    return [str(int(match)) for match in re.findall(r"\d+", text)]


def ypan_ids_from_cell(value: object) -> list[str]:
    if value is None:
        return []
    return [str(int(match)) for match in re.findall(r"(?i)YPAN\s*(\d+)", str(value))]


def load_lot_plan(
    xlsx_path: Path, reference_batch_size: int
) -> tuple[dict[tuple[str, str], SamplePlan], dict[PaperKey, int]]:
    workbook = load_workbook(xlsx_path, data_only=True, read_only=True)
    sheet_names = [
        name for name in workbook.sheetnames if re.match(r"(?i)^DAY\s+\d+$", name.strip())
    ]
    sample_to_key: dict[tuple[str, str], PaperKey] = {}
    paper_order: dict[PaperKey, int] = {}

    for sheet_name in sheet_names:
        ws = workbook[sheet_name]
        headers = [
            normalize_plan_group(value)
            for value in next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
        ]
        current_lot: int | None = None
        for row in ws.iter_rows(min_row=3, values_only=True):
            lot_cell = row[0] if len(row) > 0 else None
            paper_cell = row[1] if len(row) > 1 else None
            lot_match = LOT_RE.search(str(lot_cell)) if lot_cell is not None else None
            if lot_match:
                current_lot = int(lot_match.group(1))
            paper_match = PAPER_RE.search(str(paper_cell)) if paper_cell is not None else None
            if current_lot is None or not paper_match:
                continue

            paper_key = PaperKey(current_lot, int(paper_match.group(1)))
            if paper_key not in paper_order:
                paper_order[paper_key] = len(paper_order) + 1

            for col_idx, value in enumerate(row):
                if col_idx < 2 or col_idx >= len(headers):
                    continue
                group = headers[col_idx]
                if group is None:
                    continue

                if group == "CPAN":
                    for sample_id in numeric_ids_from_cell(value, ignore_parentheses=True):
                        sample_to_key[(group, sample_id)] = paper_key
                    for sample_id in ypan_ids_from_cell(value):
                        sample_to_key[("YPAN", sample_id)] = paper_key
                    continue

                for sample_id in numeric_ids_from_cell(value, ignore_parentheses=True):
                    sample_to_key[(group, sample_id)] = paper_key

    sample_plan: dict[tuple[str, str], SamplePlan] = {}
    for sample_key, paper_key in sample_to_key.items():
        paper_idx = paper_order[paper_key]
        sample_plan[sample_key] = SamplePlan(
            group=sample_key[0],
            sample_id=sample_key[1],
            paper_key=paper_key,
            paper_index=paper_idx,
            reference_batch=math.ceil(paper_idx / reference_batch_size),
        )
    return sample_plan, paper_order


def parse_blank_file(path: Path) -> PaperKey | None:
    match = BLANK_FILE_RE.match(path.stem)
    if match is None:
        return None
    return PaperKey(int(match.group("lot")), int(match.group("paper")))


def collect_blank_files_by_paper(root: Path) -> dict[PaperKey, list[Path]]:
    out: dict[PaperKey, list[Path]] = defaultdict(list)
    for path in readable_spectrum_files(root):
        if not any("blank" in part.lower() for part in path.parts):
            continue
        if is_average_file(path):
            continue
        paper_key = parse_blank_file(path)
        if paper_key is not None:
            out[paper_key].append(path)
    return {key: sorted(paths) for key, paths in out.items()}


def date_roots_for_controls(clinical_root: Path) -> list[Path]:
    if date_key_from_path(clinical_root) != "primary":
        return [clinical_root]
    return sorted(
        path
        for path in clinical_root.iterdir()
        if path.is_dir() and re.match(r"^20\d{6}", path.name)
    )


def collect_reference_ps_files_by_batch(
    clinical_root: Path,
    paper_order: dict[PaperKey, int],
    reference_batch_size: int,
) -> dict[int, ReferenceControl]:
    direct_by_batch: dict[int, list[Path]] = {}
    for date_root in date_roots_for_controls(clinical_root):
        date_papers = []
        blank_dir = date_root / "0. Blank"
        for path in readable_spectrum_files(blank_dir):
            paper_key = parse_blank_file(path)
            if paper_key is not None and paper_key in paper_order:
                date_papers.append(paper_key)
        if not date_papers:
            continue
        batch = math.ceil(max(paper_order[key] for key in date_papers) / reference_batch_size)
        reference_paths = collect_control_files(date_root, "ps")
        if reference_paths and batch not in direct_by_batch:
            direct_by_batch[batch] = reference_paths

    by_batch: dict[int, ReferenceControl] = {}
    last_source_batch: int | None = None
    last_paths: list[Path] = []
    max_batch = math.ceil(max(paper_order.values()) / reference_batch_size) if paper_order else 0
    for batch in range(1, max_batch + 1):
        if batch in direct_by_batch:
            last_source_batch = batch
            last_paths = direct_by_batch[batch]
        if last_source_batch is not None and last_paths:
            by_batch[batch] = ReferenceControl(
                requested_batch=batch,
                source_batch=last_source_batch,
                paths=last_paths,
            )
    return by_batch


@dataclass
class LotControlResolver:
    sample_plan: dict[tuple[str, str], SamplePlan]
    blank_files_by_paper: dict[PaperKey, list[Path]]
    reference_ps_by_batch: dict[int, ReferenceControl]

    @classmethod
    def build(
        cls, clinical_root: Path, lot_plan_xlsx: Path, reference_batch_size: int
    ) -> "LotControlResolver":
        sample_plan, paper_order = load_lot_plan(lot_plan_xlsx, reference_batch_size)
        return cls(
            sample_plan=sample_plan,
            blank_files_by_paper=collect_blank_files_by_paper(clinical_root),
            reference_ps_by_batch=collect_reference_ps_files_by_batch(
                clinical_root,
                paper_order,
                reference_batch_size,
            ),
        )

    def plan_for(self, group: str, sample_id: str) -> SamplePlan | None:
        return self.sample_plan.get((group, str(int(sample_id))))

    def blank_files_for(self, plan: SamplePlan | None) -> list[Path]:
        if plan is None:
            return []
        return self.blank_files_by_paper.get(plan.paper_key, [])

    def reference_ps_files_for(self, plan: SamplePlan | None) -> list[Path]:
        if plan is None:
            return []
        control = self.reference_ps_by_batch.get(plan.reference_batch)
        return control.paths if control is not None else []

    def reference_control_for(self, plan: SamplePlan | None) -> ReferenceControl | None:
        if plan is None:
            return None
        return self.reference_ps_by_batch.get(plan.reference_batch)


def clinical_search_roots(clinical_root: Path, date_keys: set[str] | None) -> list[Path]:
    if not date_keys:
        return [clinical_root]
    if date_key_from_path(clinical_root) in date_keys:
        return [clinical_root]
    return sorted(
        path
        for path in clinical_root.iterdir()
        if path.is_dir() and date_key_from_path(path) in date_keys
    )


def collect_matched_samples(
    primary_root: Path, clinical_root: Path, date_keys: set[str] | None = None
) -> list[dict]:
    primary_by_subject: dict[tuple[str, str], list[Path]] = defaultdict(list)
    clinical_by_subject_date: dict[tuple[str, str, str], list[Path]] = defaultdict(list)

    for path in sample_files(primary_root):
        key = parse_sample(path)
        if key is not None:
            primary_by_subject[key.subject_key].append(path)

    for root in clinical_search_roots(clinical_root, date_keys):
        for path in sample_files(root):
            key = parse_sample(path)
            if key is not None:
                date_key = date_key_from_path(path)
                if date_keys and date_key not in date_keys:
                    continue
                clinical_by_subject_date[(key.group, key.sample_id, date_key)].append(path)

    matched = []
    for (group, sample_id, date_key), clinical_paths in clinical_by_subject_date.items():
        primary_paths = primary_by_subject.get((group, sample_id), [])
        if not primary_paths:
            continue
        matched.append(
            {
                "group": group,
                "sample_id": sample_id,
                "date_key": date_key,
                "date_root": date_root_from_path(clinical_paths[0]),
                "primary_paths": sorted(primary_paths),
                "clinical_paths": sorted(clinical_paths),
            }
        )
    return sorted(matched, key=sort_key)


def read_rows(
    paths: list[Path], x_min: float, x_max: float
) -> list[tuple[Path, np.ndarray, np.ndarray]]:
    rows = []
    for path in paths:
        x, y = read_trimmed_xy(path, x_min, x_max)
        rows.append((path, x, y))
    return rows


def file_stats(
    dataset: str,
    group: str,
    sample_id: str,
    date_key: str,
    rows: list[tuple[Path, np.ndarray, np.ndarray]],
) -> list[dict]:
    out = []
    for path, x, y in rows:
        n_neg = int(np.sum(y < 0)) if len(y) else 0
        abs_y = np.abs(y)
        out.append(
            {
                "dataset": dataset,
                "group": group,
                "sample_id": sample_id,
                "date_key": date_key,
                "path": str(path),
                "n_points_after_trim": int(len(y)),
                "n_negative_after_trim": n_neg,
                "negative_fraction_after_trim": float(n_neg / len(y)) if len(y) else np.nan,
                "has_negative_after_trim": bool(n_neg > 0),
                "x_min_after_trim": float(np.min(x)) if len(x) else np.nan,
                "x_max_after_trim": float(np.max(x)) if len(x) else np.nan,
                "y_min_after_trim": float(np.min(y)) if len(y) else np.nan,
                "y_max_after_trim": float(np.max(y)) if len(y) else np.nan,
                "y_median_after_trim": float(np.median(y)) if len(y) else np.nan,
                "abs_y_p95_after_trim": float(np.percentile(abs_y, 95)) if len(y) else np.nan,
                "abs_y_max_after_trim": float(np.max(abs_y)) if len(y) else np.nan,
            }
        )
    return out


def summarize_sample(
    item: dict,
    primary_rows: list[tuple[Path, np.ndarray, np.ndarray]],
    clinical_rows: list[tuple[Path, np.ndarray, np.ndarray]],
    blank_rows: list[tuple[Path, np.ndarray, np.ndarray]],
    reference_ps_rows: list[tuple[Path, np.ndarray, np.ndarray]],
    plan: SamplePlan | None,
    reference_control: ReferenceControl | None,
) -> dict:
    primary_y = (
        np.concatenate([y for _, _, y in primary_rows if len(y)])
        if any(len(y) for _, _, y in primary_rows)
        else np.array([])
    )
    clinical_y = (
        np.concatenate([y for _, _, y in clinical_rows if len(y)])
        if any(len(y) for _, _, y in clinical_rows)
        else np.array([])
    )
    blank_y = (
        np.concatenate([y for _, _, y in blank_rows if len(y)])
        if any(len(y) for _, _, y in blank_rows)
        else np.array([])
    )
    reference_ps_y = (
        np.concatenate([y for _, _, y in reference_ps_rows if len(y)])
        if any(len(y) for _, _, y in reference_ps_rows)
        else np.array([])
    )

    def p95_abs(y: np.ndarray) -> float:
        return float(np.percentile(np.abs(y), 95)) if len(y) else np.nan

    def neg_frac(y: np.ndarray) -> float:
        return float(np.mean(y < 0)) if len(y) else np.nan

    primary_p95 = p95_abs(primary_y)
    clinical_p95 = p95_abs(clinical_y)
    ratio = clinical_p95 / primary_p95 if np.isfinite(primary_p95) and primary_p95 > 0 else np.nan
    return {
        "group": item["group"],
        "sample_id": item["sample_id"],
        "date_key": item["date_key"],
        "control_lot": plan.paper_key.lot if plan is not None else np.nan,
        "control_paper": plan.paper_key.paper if plan is not None else np.nan,
        "control_paper_index": plan.paper_index if plan is not None else np.nan,
        "control_reference_batch": plan.reference_batch if plan is not None else np.nan,
        "reference_ps_source_batch": reference_control.source_batch
        if reference_control is not None
        else np.nan,
        "reference_ps_carried_forward": bool(reference_control.carried_forward)
        if reference_control is not None
        else False,
        "n_primary_files": len(primary_rows),
        "n_clinical_files": len(clinical_rows),
        "n_blank_files": len(blank_rows),
        "n_reference_ps_files": len(reference_ps_rows),
        "n_primary_points_after_trim": int(sum(len(y) for _, _, y in primary_rows)),
        "n_clinical_points_after_trim": int(sum(len(y) for _, _, y in clinical_rows)),
        "n_blank_points_after_trim": int(sum(len(y) for _, _, y in blank_rows)),
        "n_reference_ps_points_after_trim": int(sum(len(y) for _, _, y in reference_ps_rows)),
        "primary_y_min_after_trim": float(np.min(primary_y)) if len(primary_y) else np.nan,
        "primary_y_max_after_trim": float(np.max(primary_y)) if len(primary_y) else np.nan,
        "clinical_y_min_after_trim": float(np.min(clinical_y)) if len(clinical_y) else np.nan,
        "clinical_y_max_after_trim": float(np.max(clinical_y)) if len(clinical_y) else np.nan,
        "blank_y_min_after_trim": float(np.min(blank_y)) if len(blank_y) else np.nan,
        "blank_y_max_after_trim": float(np.max(blank_y)) if len(blank_y) else np.nan,
        "reference_ps_y_min_after_trim": float(np.min(reference_ps_y))
        if len(reference_ps_y)
        else np.nan,
        "reference_ps_y_max_after_trim": float(np.max(reference_ps_y))
        if len(reference_ps_y)
        else np.nan,
        "primary_negative_fraction_after_trim": neg_frac(primary_y),
        "clinical_negative_fraction_after_trim": neg_frac(clinical_y),
        "blank_negative_fraction_after_trim": neg_frac(blank_y),
        "reference_ps_negative_fraction_after_trim": neg_frac(reference_ps_y),
        "primary_abs_y_p95_after_trim": primary_p95,
        "clinical_abs_y_p95_after_trim": clinical_p95,
        "blank_abs_y_p95_after_trim": p95_abs(blank_y),
        "reference_ps_abs_y_p95_after_trim": p95_abs(reference_ps_y),
        "clinical_to_primary_abs_y_p95_ratio_after_trim": ratio,
        "log10_clinical_to_primary_abs_y_p95_ratio_after_trim": float(np.log10(ratio))
        if np.isfinite(ratio) and ratio > 0
        else np.nan,
    }


def plot_sample_report(
    item: dict,
    primary_rows: list[tuple[Path, np.ndarray, np.ndarray]],
    clinical_rows: list[tuple[Path, np.ndarray, np.ndarray]],
    blank_rows: list[tuple[Path, np.ndarray, np.ndarray]],
    reference_ps_rows: list[tuple[Path, np.ndarray, np.ndarray]],
    out_dir: Path,
    x_min: float,
    x_max: float,
    mark_negatives: bool,
    plan: SamplePlan | None,
    reference_control: ReferenceControl | None,
    control_mode: str,
) -> Path:
    label = safe_label(item["group"], item["sample_id"], item["date_key"])
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"{label}_trim_only.png"

    fig, axes = plt.subplots(5, 1, figsize=(12, 14), sharex=True)
    if control_mode != "lot":
        blank_title = "date-level Blank: own y-axis"
        reference_title = "date-level Reference PS: own y-axis"
        control_suffix = "controls: date-level"
    elif plan is None:
        blank_title = "matched Blank: no lot-plan match"
        reference_title = "matched Reference PS: no lot-plan match"
        control_suffix = "controls: no lot-plan match"
    else:
        blank_title = f"matched Blank: {plan.paper_key.label}"
        if reference_control is not None and reference_control.carried_forward:
            reference_title = (
                f"matched Reference PS: batch {plan.reference_batch}, "
                f"using previous batch {reference_control.source_batch}"
            )
            control_suffix = (
                f"controls: {plan.paper_key.label}, Ref batch {plan.reference_batch} "
                f"from previous batch {reference_control.source_batch}"
            )
        else:
            reference_title = (
                f"matched Reference PS: batch {plan.reference_batch} after 15-paper interval"
            )
            control_suffix = f"controls: {plan.paper_key.label}, Ref batch {plan.reference_batch}"
    panels = [
        (axes[0], primary_rows, "primary raw_data: own y-axis", "#606060"),
        (axes[1], clinical_rows, "clinical reacquired: own y-axis", "#2f6f9f"),
        (axes[2], blank_rows, blank_title, "#8f8f8f"),
        (axes[3], reference_ps_rows, reference_title, "#8a5a2b"),
    ]
    for ax, rows, title, color in panels:
        for path, x, y in rows:
            ax.plot(x, y, linewidth=0.8, alpha=0.75, label=path.name, color=color)
            if mark_negatives and len(y):
                neg = y < 0
                if np.any(neg):
                    ax.scatter(x[neg], y[neg], s=3, alpha=0.30, color="#b33a3a", linewidths=0)
        ax.axhline(0, color="black", linewidth=0.7)
        ax.set_title(title, fontsize=10)
        ax.set_ylabel("CSV intensity")
        if rows and len(rows) <= 15:
            ax.legend(fontsize=5, ncol=5, loc="upper right")
        elif rows:
            ax.text(
                0.99,
                0.92,
                f"{len(rows)} files",
                ha="right",
                va="top",
                fontsize=8,
                transform=ax.transAxes,
                bbox={"facecolor": "white", "alpha": 0.75, "edgecolor": "none", "pad": 2},
            )
        else:
            ax.text(0.5, 0.5, "No files found", ha="center", va="center", transform=ax.transAxes)

    all_y = [y for _, _, y in primary_rows + clinical_rows if len(y)]
    ax = axes[4]
    for path, x, y in primary_rows:
        ax.plot(x, y, linewidth=0.75, alpha=0.65, color="#606060")
    for path, x, y in clinical_rows:
        ax.plot(x, y, linewidth=0.75, alpha=0.65, color="#2f6f9f")
    if all_y:
        y_all = np.concatenate(all_y)
        y_min = float(np.min(y_all))
        y_max = float(np.max(y_all))
        pad = max((y_max - y_min) * 0.05, 1.0)
        ax.set_ylim(y_min - pad, y_max + pad)
    ax.axhline(0, color="black", linewidth=0.7)
    ax.set_title("primary + clinical overlaid: same y-axis", fontsize=10)
    ax.set_ylabel("CSV intensity")
    ax.set_xlabel("CSV Raman shift")

    fig.suptitle(
        f"{item['group']} {item['sample_id']} | {item['date_key']} | TRIM ONLY x=[{x_min:g}, {x_max:g}] | native CSV values | {control_suffix} | not corrected",
        fontsize=11,
    )
    fig.tight_layout(rect=[0, 0, 1, 0.96])
    fig.savefig(out_path, dpi=160)
    plt.close(fig)
    return out_path


def plot_contact_sheet(
    page_items: list[dict],
    out_path: Path,
    x_min: float,
    x_max: float,
    n_cols: int,
) -> None:
    n_rows = math.ceil(len(page_items) / n_cols)
    fig, axes = plt.subplots(n_rows, n_cols, figsize=(4.4 * n_cols, 3.0 * n_rows), squeeze=False)
    for ax in axes.ravel():
        ax.axis("off")

    for ax, item in zip(axes.ravel(), page_items):
        ax.axis("on")
        primary_rows = item["primary_rows"]
        clinical_rows = item["clinical_rows"]
        for _, x, y in primary_rows:
            ax.plot(x, y, linewidth=0.45, alpha=0.55, color="#606060")
        for _, x, y in clinical_rows:
            ax.plot(x, y, linewidth=0.45, alpha=0.65, color="#2f6f9f")
        ax.axhline(0, color="black", linewidth=0.45)
        ax.set_title(f"{item['group']} {item['sample_id']} | {item['date_key']}", fontsize=8)
        ax.tick_params(axis="both", labelsize=6, length=2)
        ax.set_xlim(x_min, x_max)
    fig.suptitle(
        f"TRIM ONLY sample contact sheet | x=[{x_min:g}, {x_max:g}] | gray=primary, blue=clinical | no preprocessing",
        fontsize=12,
    )
    fig.tight_layout(rect=[0, 0, 1, 0.96])
    out_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_path, dpi=160)
    plt.close(fig)


def plot_summary_figures(sample_summary: pd.DataFrame, out_dir: Path) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    summary = sample_summary.copy()
    summary = summary.sort_values(
        ["group", "date_key", "sample_id"],
        key=lambda s: s.map(lambda v: GROUP_ORDER.index(v) if v in GROUP_ORDER else v)
        if s.name == "group"
        else s,
    )
    x = np.arange(len(summary))
    colors = [GROUP_COLORS.get(g, "#555555") for g in summary["group"]]

    fig, axes = plt.subplots(3, 1, figsize=(15, 9), sharex=True)
    axes[0].scatter(
        x,
        summary["log10_clinical_to_primary_abs_y_p95_ratio_after_trim"],
        s=14,
        c=colors,
        alpha=0.85,
    )
    axes[0].axhline(0, color="black", linewidth=0.8)
    axes[0].set_ylabel("log10 clinical / primary\nabs intensity p95")
    axes[0].set_title("Trim-only intensity scale shift by sample")

    axes[1].scatter(
        x,
        summary["primary_negative_fraction_after_trim"],
        s=10,
        c="#606060",
        alpha=0.55,
        label="primary",
    )
    axes[1].scatter(
        x,
        summary["clinical_negative_fraction_after_trim"],
        s=10,
        c="#2f6f9f",
        alpha=0.55,
        label="clinical",
    )
    axes[1].set_ylabel("negative fraction")
    axes[1].legend(fontsize=8)

    axes[2].scatter(
        x, summary["primary_abs_y_p95_after_trim"], s=10, c="#606060", alpha=0.45, label="primary"
    )
    axes[2].scatter(
        x, summary["clinical_abs_y_p95_after_trim"], s=10, c="#2f6f9f", alpha=0.45, label="clinical"
    )
    axes[2].set_yscale("log")
    axes[2].set_ylabel("abs intensity p95")
    axes[2].set_xlabel("matched sample-date index")
    axes[2].legend(fontsize=8)

    handles = [
        plt.Line2D(
            [0],
            [0],
            marker="o",
            color="w",
            markerfacecolor=GROUP_COLORS.get(g, "#555555"),
            label=g,
            markersize=6,
        )
        for g in GROUP_ORDER
        if g in set(summary["group"])
    ]
    fig.legend(
        handles=handles,
        loc="upper center",
        ncol=min(len(handles), 13),
        fontsize=8,
        bbox_to_anchor=(0.5, 0.985),
    )
    fig.tight_layout(rect=[0, 0, 1, 0.94])
    path = out_dir / "trim_only_all_samples_metric_overview.png"
    fig.savefig(path, dpi=180)
    plt.close(fig)
    paths.append(path)

    grouped = [
        summary.loc[summary["group"] == g, "log10_clinical_to_primary_abs_y_p95_ratio_after_trim"]
        .dropna()
        .to_numpy()
        for g in GROUP_ORDER
        if g in set(summary["group"])
    ]
    labels = [g for g in GROUP_ORDER if g in set(summary["group"])]
    fig, ax = plt.subplots(figsize=(12, 5))
    ax.boxplot(grouped, tick_labels=labels, showfliers=False)
    for i, values in enumerate(grouped, start=1):
        jitter = np.linspace(-0.18, 0.18, len(values)) if len(values) else []
        ax.scatter(
            np.full(len(values), i) + jitter,
            values,
            s=10,
            alpha=0.35,
            color=GROUP_COLORS.get(labels[i - 1], "#555555"),
        )
    ax.axhline(0, color="black", linewidth=0.8)
    ax.set_ylabel("log10 clinical / primary abs intensity p95")
    ax.set_title("Trim-only intensity scale shift by group")
    fig.tight_layout()
    path = out_dir / "trim_only_group_intensity_ratio_boxplot.png"
    fig.savefig(path, dpi=180)
    plt.close(fig)
    paths.append(path)
    return paths


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--primary-root", type=Path, default=DEFAULT_PRIMARY_ROOT)
    parser.add_argument("--clinical-root", type=Path, default=DEFAULT_CLINICAL_ROOT)
    parser.add_argument(
        "--date-key",
        action="append",
        default=None,
        help="Restrict clinical acquisition date, e.g. 20260416. Repeat for multiple dates.",
    )
    parser.add_argument(
        "--lot-plan-xlsx",
        type=Path,
        default=None,
        help="Excel workbook with DAY sheets mapping samples to Lot/Paper. Enables matched Blank/Reference controls.",
    )
    parser.add_argument(
        "--reference-batch-size",
        type=int,
        default=15,
        help="Number of Paper positions between Reference Material measurements.",
    )
    parser.add_argument("--trim-min", type=float, default=400.0)
    parser.add_argument("--trim-max", type=float, default=1800.0)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUT_DIR)
    parser.add_argument("--max-samples", type=int, default=None)
    parser.add_argument("--sample-reports", action=argparse.BooleanOptionalAction, default=True)
    parser.add_argument("--contact-sheets", action=argparse.BooleanOptionalAction, default=True)
    parser.add_argument("--contact-sheet-cols", type=int, default=3)
    parser.add_argument("--contact-sheet-page-size", type=int, default=12)
    parser.add_argument("--mark-negatives", action=argparse.BooleanOptionalAction, default=False)
    parser.add_argument("--overwrite", action=argparse.BooleanOptionalAction, default=False)
    args = parser.parse_args()

    date_keys = set(args.date_key) if args.date_key else None
    matched = collect_matched_samples(args.primary_root, args.clinical_root, date_keys)
    if args.max_samples is not None:
        matched = matched[: args.max_samples]
    if not matched:
        raise RuntimeError("No matched primary/clinical sample-date pairs found.")

    sample_report_dir = args.output_dir / "sample_reports"
    contact_sheet_dir = args.output_dir / "contact_sheets"
    summary_dir = args.output_dir / "summary"
    lot_control_resolver = (
        LotControlResolver.build(args.clinical_root, args.lot_plan_xlsx, args.reference_batch_size)
        if args.lot_plan_xlsx is not None
        else None
    )
    control_mode = "lot" if lot_control_resolver is not None else "date"
    file_stat_rows = []
    sample_summary_rows = []
    contact_items = []
    sample_report_paths = []
    control_row_cache: dict[
        str,
        tuple[list[tuple[Path, np.ndarray, np.ndarray]], list[tuple[Path, np.ndarray, np.ndarray]]],
    ] = {}
    blank_row_cache: dict[PaperKey, list[tuple[Path, np.ndarray, np.ndarray]]] = {}
    reference_ps_row_cache: dict[int, list[tuple[Path, np.ndarray, np.ndarray]]] = {}

    for idx, item in enumerate(matched, start=1):
        primary_rows = read_rows(item["primary_paths"], args.trim_min, args.trim_max)
        clinical_rows = read_rows(item["clinical_paths"], args.trim_min, args.trim_max)
        if lot_control_resolver is not None:
            plan = lot_control_resolver.plan_for(item["group"], item["sample_id"])
            if plan is None:
                reference_control = None
                blank_rows = []
                reference_ps_rows = []
            else:
                reference_control = lot_control_resolver.reference_control_for(plan)
                if plan.paper_key not in blank_row_cache:
                    blank_row_cache[plan.paper_key] = read_rows(
                        lot_control_resolver.blank_files_for(plan),
                        args.trim_min,
                        args.trim_max,
                    )
                if plan.reference_batch not in reference_ps_row_cache:
                    reference_ps_row_cache[plan.reference_batch] = read_rows(
                        lot_control_resolver.reference_ps_files_for(plan),
                        args.trim_min,
                        args.trim_max,
                    )
                blank_rows = blank_row_cache[plan.paper_key]
                reference_ps_rows = reference_ps_row_cache[plan.reference_batch]
        else:
            plan = None
            reference_control = None
            if item["date_key"] not in control_row_cache:
                blank_paths = collect_control_files(item["date_root"], "blank")
                reference_ps_paths = collect_control_files(item["date_root"], "ps")
                control_row_cache[item["date_key"]] = (
                    read_rows(blank_paths, args.trim_min, args.trim_max),
                    read_rows(reference_ps_paths, args.trim_min, args.trim_max),
                )
            blank_rows, reference_ps_rows = control_row_cache[item["date_key"]]
        file_stat_rows.extend(
            file_stats(
                "primary_raw_data", item["group"], item["sample_id"], item["date_key"], primary_rows
            )
        )
        file_stat_rows.extend(
            file_stats(
                "clinical_reacquired",
                item["group"],
                item["sample_id"],
                item["date_key"],
                clinical_rows,
            )
        )
        file_stat_rows.extend(
            file_stats(
                f"{control_mode}_blank",
                item["group"],
                item["sample_id"],
                item["date_key"],
                blank_rows,
            )
        )
        file_stat_rows.extend(
            file_stats(
                f"{control_mode}_reference_ps",
                item["group"],
                item["sample_id"],
                item["date_key"],
                reference_ps_rows,
            )
        )
        sample_summary_rows.append(
            summarize_sample(
                item,
                primary_rows,
                clinical_rows,
                blank_rows,
                reference_ps_rows,
                plan,
                reference_control,
            )
        )

        label = safe_label(item["group"], item["sample_id"], item["date_key"])
        if args.sample_reports:
            out_path = sample_report_dir / f"{idx:04d}_{label}_trim_only.png"
            if args.overwrite or not out_path.exists():
                plot_path = plot_sample_report(
                    item,
                    primary_rows,
                    clinical_rows,
                    blank_rows,
                    reference_ps_rows,
                    sample_report_dir,
                    args.trim_min,
                    args.trim_max,
                    args.mark_negatives,
                    plan,
                    reference_control,
                    control_mode,
                )
                plot_path.rename(out_path)
            sample_report_paths.append(out_path)

        if args.contact_sheets:
            contact_item = dict(item)
            contact_item["primary_rows"] = primary_rows
            contact_item["clinical_rows"] = clinical_rows
            contact_item["blank_rows"] = blank_rows
            contact_item["reference_ps_rows"] = reference_ps_rows
            contact_items.append(contact_item)

        if idx % 100 == 0:
            print(f"Processed {idx}/{len(matched)} matched sample-date pairs")

    args.output_dir.mkdir(parents=True, exist_ok=True)
    file_stats_path = args.output_dir / "trim_only_file_stats_all_samples.csv"
    sample_summary_path = args.output_dir / "trim_only_sample_summary_all_samples.csv"
    pd.DataFrame(file_stat_rows).to_csv(file_stats_path, index=False, encoding="utf-8-sig")
    sample_summary = pd.DataFrame(sample_summary_rows)
    sample_summary.to_csv(sample_summary_path, index=False, encoding="utf-8-sig")

    summary_paths = plot_summary_figures(sample_summary, summary_dir)
    contact_paths = []
    if args.contact_sheets:
        page_size = args.contact_sheet_page_size
        for page_idx in range(0, len(contact_items), page_size):
            page = contact_items[page_idx : page_idx + page_size]
            out_path = (
                contact_sheet_dir / f"trim_only_contact_sheet_{page_idx // page_size + 1:03d}.png"
            )
            plot_contact_sheet(
                page, out_path, args.trim_min, args.trim_max, args.contact_sheet_cols
            )
            contact_paths.append(out_path)

    manifest_path = args.output_dir / "trim_only_outputs_manifest.txt"
    with manifest_path.open("w", encoding="utf-8") as fh:
        fh.write(f"matched_sample_date_pairs={len(matched)}\n")
        fh.write(f"control_mode={control_mode}\n")
        if args.lot_plan_xlsx is not None:
            fh.write(f"lot_plan_xlsx={args.lot_plan_xlsx}\n")
            fh.write(f"reference_batch_size={args.reference_batch_size}\n")
        fh.write(f"sample_reports={len(sample_report_paths)}\n")
        fh.write(f"contact_sheets={len(contact_paths)}\n")
        fh.write(f"file_stats={file_stats_path}\n")
        fh.write(f"sample_summary={sample_summary_path}\n")
        for path in summary_paths:
            fh.write(f"summary_figure={path}\n")
        for path in contact_paths:
            fh.write(f"contact_sheet={path}\n")
        for path in sample_report_paths:
            fh.write(f"sample_report={path}\n")

    print(f"Matched sample-date pairs: {len(matched)}")
    print(f"Control mode: {control_mode}")
    if args.lot_plan_xlsx is not None:
        print(f"Lot plan: {args.lot_plan_xlsx}")
    print(f"Sample reports: {len(sample_report_paths)}")
    print(f"Contact sheets: {len(contact_paths)}")
    print(f"File stats: {file_stats_path}")
    print(f"Sample summary: {sample_summary_path}")
    print(f"Manifest: {manifest_path}")
    for path in summary_paths:
        print(f"Summary figure: {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
