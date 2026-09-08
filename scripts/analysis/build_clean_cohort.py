#!/usr/bin/env python3
"""Build a clean retrospective clinical/SERS cohort for model training.

Clean cohort definition:
  - study_design == retrospective
  - is_pre_effective == 1
  - history_other_cancer == 0
  - groups used by the 7-cancer + control screening model
  - spectra available in the selected processed_spectra.csv
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_CLINICAL = PROJECT_ROOT / "data" / "clinical_data" / "전체환자_임상정보_정규화.xlsx"
DEFAULT_SPECTRA = PROJECT_ROOT / "results" / "processed_spectra.csv"
DEFAULT_OUT = PROJECT_ROOT / "results" / "clean_cohort_20260605"

TARGET_SOURCE_GROUPS = {
    "PRO",
    "BRE",
    "OVA",
    "LUN",
    "CRC",
    "CPAN",
    "BLC",
    "NOR",
    "DIA",
    "HBP",
    "H.D.",
}
GROUP_DISPLAY = {
    "PRO": "prostate",
    "BRE": "breast",
    "OVA": "ovarian",
    "LUN": "lung",
    "CRC": "colorectal",
    "CPAN": "pancreatic",
    "BLC": "bladder",
    "NOR": "control",
    "DIA": "control",
    "HBP": "control",
    "H.D.": "control",
}
MODEL_GROUP = {
    "CPAN": "PAN",
    "PRO": "PRO",
    "BRE": "BRE",
    "OVA": "OVA",
    "LUN": "LUN",
    "CRC": "CRC",
    "BLC": "BLC",
    "NOR": "NOR",
    "DIA": "DIA",
    "HBP": "HBP",
    "H.D.": "H.D.",
}


def parse_solum_label(value: object) -> tuple[str, str | None]:
    text = str(value or "").strip()
    match = re.match(r"^\s*([A-Za-z]+(?:\.?\s*[A-Za-z]+\.?)?)\s*_?\s*(\d+)", text)
    if not match:
        return text.upper(), None
    prefix = match.group(1).upper().replace(" ", "")
    if prefix in {"H.D", "H.D."}:
        prefix = "H.D."
    if prefix == "OVARIAN":
        prefix = "OVA"
    return prefix, str(int(match.group(2)))


def load_clinical(path: Path) -> pd.DataFrame:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Sheet1"]
    headers = [cell.value for cell in ws[1]]
    idx = {h: i for i, h in enumerate(headers) if h is not None}
    required = [
        "group",
        "solum_label",
        "patient_code",
        "study_design",
        "hospital_code",
        "history_other_cancer",
        "is_pre_effective",
        "has_cancer_hx",
        "has_different_cancer",
        "different_cancer_sites",
    ]
    missing = [col for col in required if col not in idx]
    if missing:
        raise ValueError(f"Missing clinical columns: {missing}")

    rows = []
    for excel_row, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        source_group, sample_id = parse_solum_label(row[idx["solum_label"]])
        if source_group not in TARGET_SOURCE_GROUPS:
            continue
        rows.append(
            {
                "excel_row": excel_row,
                "group": row[idx["group"]],
                "model_group": MODEL_GROUP[source_group],
                "source_group": source_group,
                "sample_id": sample_id,
                "cohort_subject_id": f"{source_group}_{sample_id}",
                "solum_label": row[idx["solum_label"]],
                "patient_code": row[idx["patient_code"]],
                "study_design": row[idx["study_design"]],
                "hospital_code": row[idx["hospital_code"]],
                "is_pre_effective": int(row[idx["is_pre_effective"]] or 0),
                "history_other_cancer": int(row[idx["history_other_cancer"]] or 0),
                "has_cancer_hx": int(row[idx["has_cancer_hx"]] or 0),
                "has_different_cancer": int(row[idx["has_different_cancer"]] or 0),
                "different_cancer_sites": row[idx["different_cancer_sites"]],
            }
        )
    return pd.DataFrame(rows)


def build_outputs(clinical_path: Path, spectra_path: Path, out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    clinical = load_clinical(clinical_path)

    clinical["eligible_clean_clinical"] = (
        (clinical["study_design"] == "retrospective")
        & (clinical["is_pre_effective"] == 1)
        & (clinical["history_other_cancer"] == 0)
    )
    clean_clinical = clinical[clinical["eligible_clean_clinical"]].copy()

    spectra = pd.read_csv(spectra_path)
    spectra["source_group"] = spectra["group"].astype(str).str.upper()
    spectra["sample_id_str"] = spectra["sample_id"].astype(str).str.replace(r"\.0$", "", regex=True)

    clean_keys = set(zip(clean_clinical["source_group"], clean_clinical["sample_id"]))
    spectra_keys = set(zip(spectra["source_group"], spectra["sample_id_str"]))
    matched_keys = clean_keys & spectra_keys

    clean_clinical["has_spectrum"] = clean_clinical.apply(
        lambda r: (r["source_group"], r["sample_id"]) in spectra_keys,
        axis=1,
    )
    manifest = clean_clinical[clean_clinical["has_spectrum"]].copy()
    missing = clean_clinical[~clean_clinical["has_spectrum"]].copy()

    subset_mask = spectra.apply(
        lambda r: (r["source_group"], r["sample_id_str"]) in matched_keys,
        axis=1,
    )
    clean_spectra = spectra.loc[subset_mask].drop(columns=["source_group", "sample_id_str"])

    flow_rows = [
        {
            "step": "all_target_groups",
            "n_subjects": clinical[["source_group", "sample_id"]].drop_duplicates().shape[0],
            "n_cancer": int((clinical["group"] != "control").sum()),
            "n_control": int((clinical["group"] == "control").sum()),
        },
        {
            "step": "retrospective_only",
            "n_subjects": int((clinical["study_design"] == "retrospective").sum()),
            "n_cancer": int(
                (
                    (clinical["study_design"] == "retrospective") & (clinical["group"] != "control")
                ).sum()
            ),
            "n_control": int(
                (
                    (clinical["study_design"] == "retrospective") & (clinical["group"] == "control")
                ).sum()
            ),
        },
        {
            "step": "retrospective_and_pre_effective",
            "n_subjects": int(
                (
                    (clinical["study_design"] == "retrospective")
                    & (clinical["is_pre_effective"] == 1)
                ).sum()
            ),
            "n_cancer": int(
                (
                    (clinical["study_design"] == "retrospective")
                    & (clinical["is_pre_effective"] == 1)
                    & (clinical["group"] != "control")
                ).sum()
            ),
            "n_control": int(
                (
                    (clinical["study_design"] == "retrospective")
                    & (clinical["is_pre_effective"] == 1)
                    & (clinical["group"] == "control")
                ).sum()
            ),
        },
        {
            "step": "clean_clinical_no_other_cancer_history",
            "n_subjects": len(clean_clinical),
            "n_cancer": int((clean_clinical["group"] != "control").sum()),
            "n_control": int((clean_clinical["group"] == "control").sum()),
        },
        {
            "step": "clean_clinical_with_spectrum",
            "n_subjects": len(manifest),
            "n_cancer": int((manifest["group"] != "control").sum()),
            "n_control": int((manifest["group"] == "control").sum()),
        },
    ]
    flow = pd.DataFrame(flow_rows)

    counts = (
        manifest.groupby(["group", "source_group", "model_group", "hospital_code"], dropna=False)
        .agg(n_subjects=("sample_id", "nunique"))
        .reset_index()
        .sort_values(["group", "source_group", "hospital_code"])
    )
    spectra_counts = (
        clean_spectra.groupby("group", dropna=False)
        .agg(n_spectra=("sample_id", "size"), n_subjects=("sample_id", "nunique"))
        .reset_index()
        .sort_values("group")
    )

    manifest.to_csv(out_dir / "clean_cohort_manifest.csv", index=False, encoding="utf-8-sig")
    missing.to_csv(
        out_dir / "clean_clinical_missing_spectra.csv", index=False, encoding="utf-8-sig"
    )
    counts.to_csv(out_dir / "clean_cohort_counts.csv", index=False, encoding="utf-8-sig")
    spectra_counts.to_csv(out_dir / "clean_spectra_counts.csv", index=False, encoding="utf-8-sig")
    flow.to_csv(out_dir / "clean_cohort_filter_flow.csv", index=False, encoding="utf-8-sig")
    clean_spectra.to_csv(out_dir / "clean_processed_spectra.csv", index=False, encoding="utf-8-sig")

    print(f"Output: {out_dir}")
    print(flow.to_string(index=False))
    print("\nClean manifest by group:")
    group_counts = (
        manifest.groupby("group")
        .apply(lambda g: g[["source_group", "sample_id"]].drop_duplicates().shape[0])
        .sort_index()
    )
    print(group_counts.to_string())
    print("\nClean spectra by source group:")
    print(spectra_counts.to_string(index=False))
    if len(missing):
        print("\nClinical clean subjects missing spectra:")
        print(
            missing[["source_group", "sample_id", "solum_label", "patient_code"]].to_string(
                index=False
            )
        )


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--clinical", type=Path, default=DEFAULT_CLINICAL)
    parser.add_argument("--spectra", type=Path, default=DEFAULT_SPECTRA)
    parser.add_argument("--out-dir", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()
    build_outputs(args.clinical, args.spectra, args.out_dir)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
